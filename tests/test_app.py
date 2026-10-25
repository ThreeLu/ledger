"""端到端测试：真浏览器打开账本网页，连本地的假 GitHub（tests/fake_github.py），把主要功能走一遍。

    pip install playwright openpyxl && python -m playwright install chromium
    python tests/test_app.py            # 全部
    python tests/test_app.py 转账 校准    # 只跑名字里含这些字的步骤（前面的「开始记账」总会跑）

汇率接口用假数据（1 美元 = 7 元），不联网、不需要令牌。失败时截图在 tests/artifacts/。
"""

import json
import re
import sys
import threading
import traceback
from datetime import date, timedelta
from functools import partial
from urllib.parse import quote
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from playwright.sync_api import expect, sync_playwright

sys.path.insert(0, str(Path(__file__).parent))
from fake_github import FakeRepo, serve  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
ART = ROOT / "tests" / "artifacts"
APP_PORT, API_PORT = 8775, 8776
URL = f"http://127.0.0.1:{APP_PORT}/"
API = f"http://127.0.0.1:{API_PORT}"
REPO = "test/finance-data"
TODAY = date.today().isoformat()

STEPS = []
LAST_AI = []  # 最近一次发给 DeepSeek 的请求（检查有没有发不该发的东西）


def step(name):
    def wrap(fn):
        STEPS.append((name, fn))
        return fn
    return wrap


class Ctx:
    def __init__(self, page, repo, inventory=None):
        self.page, self.repo, self.inventory = page, repo, inventory
        self.prompt = ""  # 下一次 prompt() 弹窗填什么

    def data(self):
        # 修改是先存手机、后台上传的：等待上传队列清空再读仓库
        self.page.wait_for_function(
            "() => { try { const q = JSON.parse(localStorage.getItem('ledger-queue')); return !q || !q.items.length; } catch { return true; } }",
            timeout=15000)
        return json.loads(self.repo.read("finance.json"))

    def tx(self):
        return self.data()["tx"]

    def go(self, hash_):
        same = self.page.url == URL + hash_
        self.page.goto(URL + hash_)
        if same:
            self.page.reload()

    def recent(self, text=""):
        """首页「最近记的」：不给字看最近 8 笔，给字就点放大镜在全部时间里搜"""
        self.go("#/")
        box = self.page.locator(".recent-search")
        if bool(text) != box.is_visible():
            self.page.locator(".recent-btn").click()
        if text:
            box.get_by_label("搜索记过的账").fill(text)
        return self.page.locator(".recent-list .tx")

    def balance(self, acc):
        d = self.data()
        a = next(x for x in d["accounts"] if x["id"] == acc)
        b = a["opening"]
        for t in d["tx"]:
            if t.get("account") == acc:
                b += {"income": 1, "repay": 1, "adjust": 1, "expense": -1, "transfer": -1, "advance": -1, "payback": -1}.get(t["type"], 0) * t["amount"]
            if t["type"] == "transfer" and t.get("to") == acc:
                b += t.get("toAmount", t["amount"])
        return round(b, 2)

    def wait_saved(self, n):
        """等仓库里的流水数变成 n（保存是异步的）"""
        for _ in range(60):
            if len(self.tx()) == n:
                return
            self.page.wait_for_timeout(200)
        raise AssertionError(f"流水数是 {len(self.tx())}，不是 {n}")

    def add(self, amount, *, kind="支出", cat=None, acc=None, to=None, note=None):
        p = self.page
        self.go("#/add")
        p.locator(".segmented").get_by_role("button", name=kind, exact=True).click()
        p.get_by_label("金额", exact=True).fill(str(amount))
        if cat:
            p.get_by_role("button", name=cat, exact=True).first.click()  # 「最近用过」里可能也有
        if acc:
            group = {"支出": "账户", "收入": "账户", "转账": "转出账户"}[kind]
            p.get_by_role("group", name=group).get_by_role("button", name=acc, exact=True).click()
        if to:
            p.get_by_role("group", name="转入账户").get_by_role("button", name=to, exact=True).click()
        if note:
            p.get_by_label("备注").fill(note)
        n = len(self.tx())
        p.get_by_role("button", name="记好了").click()
        self.wait_saved(n + 1)
        p.wait_for_function("location.hash === '#/'")


# ---------- 测试步骤 ----------

@step("开始记账：用物品档案的令牌连上，填期初余额，建好账本")
def _(c):
    p = c.page
    p.goto(URL)
    p.evaluate(f"""() => {{
        localStorage.clear();
        localStorage.setItem('ledger-api-base', '{API}');
        localStorage.setItem('inventory-settings', JSON.stringify({{ repo: 'x/inventory-data', token: 'test-token' }}));
        localStorage.setItem('ledger-settings', JSON.stringify({{ repo: '{REPO}' }}));
    }}""")
    p.goto(URL)
    expect(p.get_by_role("heading", name="开始记账")).to_be_visible()
    for name, v in [("存钱卡余额", "10000"), ("生活费卡余额", "500"), ("微信余额", "1000"), ("校园卡余额", "100")]:
        p.get_by_label(name).fill(v)
    p.get_by_role("button", name="开始记账").click()
    expect(p.locator(".summary-title")).to_be_visible()
    d = c.data()
    assert {a["id"]: a["opening"] for a in d["accounts"]}["a-save"] == 10000, d["accounts"]
    assert d["openingDate"] == TODAY
    # 个人设置在私有仓库里改（这里用编的数字模拟）：美元账户、固定扣费、收入计划、预算
    d["accounts"].append({"id": "a-usd", "name": "Apple ID", "currency": "USD", "opening": 150, "note": ""})
    d["presets"].append({"name": "充值 Apple ID", "from": "a-live", "to": "a-usd"})
    d["recurring"] = [{"id": "r-claude", "name": "订阅甲", "amount": 20, "account": "a-usd", "category": "c-ai", "day": 6, "since": TODAY},
                      {"id": "r-gpt", "name": "订阅乙", "amount": 100, "account": "a-usd", "category": "c-ai", "day": 29, "since": TODAY}]
    d["incomePlan"] = [{"name": "收入甲", "amount": 4000, "when": "每月 15 号", "use": "过这个月"}, {"name": "收入乙", "amount": 4000, "use": "全部存下"}]
    d["settings"].update(periodStartDay=15, expectedIncome=8000, emergencyFloor=8000, summerMonths=[7, 8], sideIncomeSave=0.7, payNote="测试：支付宝记在生活费卡")
    d["budget"] = {"food": 1900, "daily": 700, "free": 300, "sub": 900}
    expect(p.locator(".health-tile", has_text="安全垫")).to_be_visible()
    # 固定扣费先推到很远以后，免得今天刚好是扣费日、打乱后面的笔数（扣费单独测）
    for r in d["recurring"]:
        r["since"] = "2999-01-01"
    d["tx"] = []
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    p.reload()
    # 等重画好；不看「一切正常」那一行，月底前后 Apple ID 充值提醒会把它换掉
    expect(p.locator(".spend-left")).to_be_visible()


@step("记支出：金额、类别、账户；没有快捷记账、导入小票")
def _(c):
    p = c.page
    c.go("#/add")
    expect(p.get_by_text("测试：支付宝记在生活费卡")).to_be_visible()
    expect(p.get_by_text("存成快捷")).to_have_count(0)
    expect(p.get_by_text("有小票？")).to_have_count(0)
    c.add(15, cat="午餐", acc="校园卡", note="食堂午饭")
    expect(p.locator(".tx", has_text="午餐").first).to_be_visible()
    d = c.data()
    t = d["tx"][-1]
    assert (t["type"], t["amount"], t["category"], t["account"], t["note"]) == ("expense", 15, "c-lunch", "a-campus", "食堂午饭"), t
    assert "quick" not in d
    c.add(15, cat="午餐", acc="校园卡", note="食堂午饭")
    assert c.balance("a-campus") == 70
    # 下次默认账户是上次用的
    c.go("#/add")
    expect(p.get_by_role("group", name="账户").get_by_role("button", name="校园卡")).to_have_attribute("aria-pressed", "true")


@step("记收入：生活费 4000 到存钱卡")
def _(c):
    c.add(4000, kind="收入", cat="生活费", acc="存钱卡")
    assert c.balance("a-save") == 14000


@step("转账：常用的「充值校园卡」带手续费；充值 Apple ID 人民币换美元")
def _(c):
    p = c.page
    c.go("#/add")
    p.locator(".segmented").get_by_role("button", name="转账").click()
    p.get_by_role("button", name="充值校园卡").click()
    p.get_by_label("金额", exact=True).fill("200")
    p.get_by_label("手续费").fill("1")
    n = len(c.tx())
    p.get_by_role("button", name="记好了").click()
    c.wait_saved(n + 2)
    assert c.balance("a-campus") == 270 and c.balance("a-live") == 299, (c.balance("a-campus"), c.balance("a-live"))
    fee = c.tx()[-1]
    assert fee["category"] == "c-fee" and fee["amount"] == 1
    p.wait_for_function("location.hash === '#/'")
    expect(p.get_by_text("另有 ¥1 是手续费的钱，不占预算")).to_be_visible()
    # 充值 Apple ID：140 元按 1:7 估成 20 美元
    c.go("#/add")
    p.locator(".segmented").get_by_role("button", name="转账").click()
    p.get_by_role("button", name="充值 Apple ID").click()
    p.get_by_label("金额", exact=True).fill("140")
    expect(p.get_by_label("到账金额")).to_have_value("20")
    p.get_by_label("到账金额").fill("19.5")
    p.get_by_role("button", name="记好了").click()
    c.wait_saved(n + 3)
    assert c.balance("a-usd") == 169.5, c.balance("a-usd")
    # 美元账户的花销按汇率折成人民币
    c.add(2, cat="软件订阅", acc="Apple ID")
    assert c.tx()[-1]["cny"] == 14


@step("最近记的：首页最近 8 笔；搜出来改一笔、删一笔")
def _(c):
    p = c.page
    expect(c.recent()).to_have_count(min(8, len(c.tx())))
    c.recent("生活费").filter(has_text=re.compile(r"^生活费\d")).click()
    expect(p.get_by_role("heading", name="改一笔")).to_be_visible()
    p.get_by_label("金额", exact=True).fill("4100")
    p.get_by_role("button", name="保存").click()
    p.wait_for_function("location.hash === '#/' || location.hash === ''")
    expect(p.locator(".recent-search input")).to_have_value("生活费")  # 改完回来还是搜索结果
    assert c.balance("a-save") == 14100
    n = len(c.tx())
    c.recent("软件订阅").first.click()
    p.get_by_role("button", name="删除").click()
    c.wait_saved(n - 1)
    # 关掉放大镜回到最近 8 笔
    expect(c.recent()).to_have_count(min(8, n - 1))
    expect(p.locator(".recent-search")).to_be_hidden()


@step("账户：总资产、校准（差额记对账差额）")
def _(c):
    p = c.page
    c.go("#/accounts")
    expect(p.locator(".big-num")).to_be_visible()
    p.locator(".cell", has_text="生活费卡").click()
    p.get_by_role("button", name="校准").click()
    p.get_by_label("实际余额").fill("480")
    n = len(c.tx())
    p.locator(".sheet").get_by_role("button", name="校准").click()
    c.wait_saved(n + 1)
    t = c.tx()[-1]
    # 生活费卡：500 − 充校园卡 200 − 手续费 1 − 充 Apple ID 140 = 159，实际 480 → 差额 +321
    assert t["type"] == "adjust" and t["amount"] == 321 and c.balance("a-live") == 480, t
    expect(p.locator(".tx", has_text="对账差额")).to_be_visible()


@step("首页：健康指标能点开看解释；应急钱跌破底线变红并说该做什么")
def _(c):
    p = c.page
    c.go("#/")
    p.locator(".health-tile", has_text="安全垫").click()
    sheet = p.locator(".sheet")
    for t in ("这是什么", "为什么重要", "你现在"):
        expect(sheet).to_contain_text(t)
    sheet.get_by_role("button", name="知道了").click()
    p.get_by_role("button", name="怎么用").click()
    expect(p.locator(".sheet")).to_contain_text("健康指标")
    p.locator(".sheet").get_by_role("button", name="知道了").click()
    # 花钱方式里「应急钱」那条规矩点开改底线：调到 99999 → 红
    def set_floor(v):
        c.go("#/rules")
        rule = p.locator(".rule-item", has_text="应急钱")
        rule.locator("summary").click()
        rule.get_by_role("button", name="改应急钱底线").click()
        p.locator(".sheet").get_by_label("应急钱底线").fill(v)
        p.locator(".sheet").get_by_role("button", name="保存").click()
        expect(rule).to_contain_text("¥" + format(int(v), ","))
    set_floor("99999")
    assert c.data()["settings"]["emergencyFloor"] == 99999
    c.go("#/")
    expect(p.locator(".todo-item.bad", has_text="应急钱")).to_contain_text("下一笔收入会先把它补上")
    set_floor("8000")

    # 每月预算在花钱方式最上面：点「改」就地改，保存后还在这一页
    def set_budget(**kw):
        c.go("#/rules")
        p.locator(".budget-plan").get_by_role("button", name="改").click()
        for label, v in kw.items():
            p.locator(".budget-plan").get_by_label(label).fill(v)
        p.locator(".budget-plan").get_by_role("button", name="保存").click()
        expect(p.locator(".budget-plan input")).to_have_count(0)
        assert p.evaluate("location.hash") == "#/rules"
    set_budget(吃饭="2000")
    c.go("#/")
    expect(p.locator(".todo-item.bad")).to_have_count(0)
    assert c.data()["budget"]["food"] == 2000
    # 形象预算：单独一组，算进生活预算
    set_budget(形象="150")
    expect(p.locator(".budget-plan")).to_contain_text("形象")
    assert c.data()["budget"]["look"] == 150
    c.add(66, cat="护肤", acc="生活费卡", note="编的洗面奶")
    assert c.tx()[-1]["category"] == "c-skin"
    c.go("#/rules")
    p.locator(".rule-item summary", has_text="形象").click()
    expect(p.get_by_text("买护肤品、化妆品、香水这些")).to_be_visible()
    # 「改」了又取消：不改
    p.locator(".budget-plan").get_by_role("button", name="改").click()
    p.locator(".budget-plan").get_by_label("吃饭").fill("1")
    p.locator(".budget-plan").get_by_role("button", name="取消").click()
    assert c.data()["budget"]["food"] == 2000
    set_budget(形象="0")


@step("固定扣费：到日子自动记一笔（美元按汇率折算），不会重复记")
def _(c):
    p = c.page
    d = c.data()
    past = (date.today() - timedelta(days=40)).isoformat()
    for r in d["recurring"]:
        if r["id"] == "r-claude":
            r["since"] = past
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    n = len(d["tx"])
    # 从 40 天前到今天，每月 6 号扣一次：一般是 1～2 次
    days = [date.today() - timedelta(days=k) for k in range(41)]
    expected = sum(1 for x in days if x.day == 6)
    p.reload()
    c.wait_saved(n + expected)
    auto = [t for t in c.tx() if t.get("auto") == "r-claude"]
    assert auto and all(t["amount"] == 20 and t["cny"] == 140 for t in auto), auto
    p.reload()
    p.wait_for_timeout(1500)
    assert len([t for t in c.tx() if t.get("auto") == "r-claude"]) == len(auto), "重复记了"


@step("花钱方式：预算在最上面，规矩折起来，一年能存多少；更多里只有一项")
def _(c):
    p = c.page
    c.go("#/more")
    expect(p.get_by_role("link", name=re.compile("^预算"))).to_have_count(0)
    p.get_by_role("link", name=re.compile("^花钱方式")).click()
    expect(p.locator(".budget-plan")).to_contain_text("每月预算")
    expect(p.locator(".rule-item summary", has_text="先存后花")).to_be_visible()
    expect(p.get_by_text("钱一到账，先把要存的留下")).to_be_hidden()  # 折着
    rule = lambda t: p.locator(".rule-item", has_text=t)  # noqa: E731
    for t in ("2 笔收入", "7、8 月没有收入", "兼职的钱"):
        rule(t).locator("summary").click()
    expect(rule("2 笔收入")).to_contain_text("收入乙 ¥4,000：全部存下")
    expect(rule("2 笔收入").locator(".set-row")).to_contain_text("每月正常收入¥8,000")
    expect(rule("兼职的钱")).to_contain_text("70% 存下")
    p.screenshot(path=ART / "rules.png", full_page=True)
    expect(p.locator(".rule", has_text="一年下来")).to_contain_text("能存约 ¥33,200，储蓄率 42%")  # 前面把吃饭预算改成了 2000

def person(c, name):
    return next(p for p in c.data()["people"] if p["name"] == name)


def net(c, pid):
    """和这个人之间：正数他欠我，负数我欠他（和网页 personStatus 一样算）"""
    n = 0
    for t in c.tx():
        if t.get("person") != pid:
            continue
        v = t.get("cny", t["amount"])
        if t["type"] in ("advance", "payback"):
            n += v
        elif t["type"] == "repay" or (t["type"] == "expense" and not t.get("account")):
            n -= v
    return round(n, 2)


@step("AA：先付 200，选两个新同学，我那份算花销、其他记成欠我；别人帮我付；人情账还钱")
def _(c):
    p = c.page
    c.go("#/add")
    p.get_by_label("金额", exact=True).fill("200")
    p.get_by_role("button", name="出去吃", exact=True).first.click()
    p.get_by_role("group", name="账户").get_by_role("button", name="微信", exact=True).click()
    p.get_by_role("group", name="和别人有关").get_by_role("button", name="AA / 帮人付").click()
    for name in ("小甲", "小乙"):
        p.get_by_label("和谁 AA").fill(name)
        p.get_by_role("option", name=f"＋ 新加「{name}」").click()
    expect(p.locator(".picker-token")).to_have_count(2)
    expect(p.get_by_label("我那份")).to_have_value("66.67")
    p.get_by_label("我那份").fill("50")
    expect(p.locator(".split-hint")).to_contain_text("其他 2 人各约 ¥75")
    n = len(c.tx())
    p.get_by_role("button", name="记好了").click()
    c.wait_saved(n + 3)
    new = c.tx()[n:]
    assert sorted((t["type"], t["amount"]) for t in new) == [("advance", 75), ("advance", 75), ("expense", 50)], new
    assert len({t["group"] for t in new}) == 1
    a, b = person(c, "小甲"), person(c, "小乙")
    assert net(c, a["id"]) == 75 and net(c, b["id"]) == 75
    # 小乙帮我付了 30：算我的花销，账户不动
    wechat = c.balance("a-wechat")
    c.go("#/add")
    p.get_by_label("金额", exact=True).fill("30")
    p.get_by_role("button", name="午餐", exact=True).first.click()
    p.get_by_role("group", name="和别人有关").get_by_role("button", name="别人帮我付的").click()
    expect(p.get_by_role("group", name="账户")).to_have_count(0)
    p.get_by_label("谁帮我付的").fill("乙")
    expect(p.locator(".picker-list")).not_to_contain_text("null")
    p.get_by_role("option", name="小乙").click()
    expect(p.locator(".picker-token")).to_have_text("小乙×")
    p.get_by_role("button", name="记好了").click()
    c.wait_saved(n + 4)
    t = c.tx()[-1]
    assert t["account"] is None and t["person"] == b["id"] and c.balance("a-wechat") == wechat, t
    assert net(c, b["id"]) == 45
    # 人情账：小甲还我 75 到生活费卡（当初是微信付的，没关系）
    c.go("#/more")
    p.get_by_role("link", name="人情账").click()
    expect(p.locator(".cell", has_text="小乙")).to_contain_text("¥45")
    p.locator(".cell", has_text="小甲").click()
    expect(p.locator(".big-num")).to_have_text("¥75")
    p.get_by_role("button", name="他还我钱").click()
    expect(p.locator(".sheet").get_by_label("金额")).to_have_value("75")
    p.locator(".sheet").get_by_role("group", name="钱到了哪个账户").get_by_role("button", name="生活费卡").click()
    live = c.balance("a-live")
    p.locator(".sheet").get_by_role("button", name="记好了").click()
    c.wait_saved(n + 5)
    assert net(c, a["id"]) == 0 and c.balance("a-live") == round(live + 75, 2)
    expect(p.locator(".big-num")).to_have_text("¥0")
    # 首页：待收回
    c.go("#/")
    expect(p.locator(".health-tile", has_text="待收回")).to_contain_text("¥45")
    # 删 AA 的那笔花销，连同记给别人的一起删
    c.recent("出去吃").filter(has_text=re.compile(r"^出去吃")).click()
    p.get_by_role("button", name="删除").click()
    c.wait_saved(n + 2)
    assert not any(t.get("group") == new[0]["group"] for t in c.tx())


@step("垫付报销：建一件事、垫两笔、传 PDF 和照片、勾已提交、分两次报销、结清（报不回的算出差自付）")
def _(c):
    p = c.page
    c.go("#/claims")
    p.get_by_role("button", name="新的垫付").click()
    p.get_by_label("这件事").fill("测试开会")
    p.locator(".sheet").get_by_role("button", name="建好").click()
    p.wait_for_function("location.hash.startsWith('#/claim/')")
    cid = c.data()["claims"][-1]["id"]
    save_before, n = c.balance("a-save"), len(c.tx())
    for amount in ("600", "400"):
        p.get_by_role("button", name="垫一笔").click()
        p.locator(".sheet").get_by_label("金额").fill(amount)
        p.locator(".sheet").get_by_role("group", name="从哪个账户付的").get_by_role("button", name="存钱卡").click()
        p.locator(".sheet").get_by_role("button", name="记好了").click()
        expect(p.locator(".sheet")).to_have_count(0)
    c.wait_saved(n + 2)
    assert c.balance("a-save") == round(save_before - 1000, 2)
    expect(p.locator(".big-num")).to_have_text("¥1,000")
    # 不算花销：这个预算月的「不占预算」里没有它
    # 发票
    pdf = ART / "invoice.pdf"
    pdf.write_bytes(b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n")
    img = ART / "ticket.png"
    import struct, zlib
    raw = b"".join(b"\x00" + b"\x80\x80\x80" * 8 for _ in range(8))
    chunk = lambda t, d: struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d))  # noqa: E731
    img.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 8, 8, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))
    p.get_by_label("上传发票").set_input_files([str(pdf), str(img)])
    expect(p.get_by_text("已提交 0 / 2")).to_be_visible()
    docs = next(x for x in c.data()["claims"] if x["id"] == cid)["docs"]
    assert [x["kind"] for x in docs] == ["pdf", "image"] and c.repo.read(docs[0]["file"]).startswith(b"%PDF"), docs
    p.get_by_label("invoice.pdf 已提交").check()
    expect(p.get_by_text("已提交 1 / 2")).to_be_visible()
    # 报销分两次到账
    for amount, acc in (("500", "生活费卡"), ("300", "微信")):
        p.get_by_role("button", name="报销到账").click()
        p.locator(".sheet").get_by_label("金额").fill(amount)
        p.locator(".sheet").get_by_role("group", name="打到哪个账户").get_by_role("button", name=acc).click()
        p.locator(".sheet").get_by_role("button", name="记好了").click()
        expect(p.locator(".sheet")).to_have_count(0)
    c.wait_saved(n + 4)
    expect(p.locator(".big-num")).to_have_text("¥200")
    p.get_by_role("button", name="结清").click()
    expect(p.locator(".sheet")).to_contain_text("还差 ¥200 报不回来")
    p.locator(".sheet").get_by_role("button", name="结清").click()
    c.wait_saved(n + 5)
    w = c.tx()[-1]
    assert (w["type"], w["amount"], w["category"], w.get("account")) == ("writeoff", 200, "c-trip", None), w
    assert next(x for x in c.data()["claims"] if x["id"] == cid)["status"] == "settled"
    c.go("#/")
    expect(p.get_by_text("出差自付").first).to_be_visible()
    # 删一张发票
    c.go(f"#/claim/{cid}")
    p.locator(".manage-row", has_text="ticket.png").get_by_role("button", name="删除").click()
    expect(p.get_by_text("已提交 1 / 1")).to_be_visible()
    assert c.repo.read(docs[1]["file"]) is None


@step("对账：随时对、和卡无关；挪了没记的只改余额，少了算漏记的花销，多了直接进存款；14 天没对提醒")
def _(c):
    p = c.page
    d = c.data()
    d["openingDate"] = "2020-01-01"  # 假装早就开始记账、一直没对过
    d["lastReconciled"] = ""
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    c.go("#/")
    p.reload()
    expect(p.locator(".health-tile", has_text="对账")).to_be_visible()
    c.go("#/reconcile")
    expect(p.locator(".card-check")).to_have_count(0)                   # 不再算两张卡之间转多少
    expect(p.locator(".rec-top-card")).to_contain_text("网站上一共")
    live, wechat = c.balance("a-live"), c.balance("a-wechat")
    # 一边少一边多：挪了没记
    p.get_by_label("生活费卡实际余额").fill(str(round(live - 30, 2)))
    p.get_by_label("微信实际余额").fill(str(round(wechat + 30, 2)))
    expect(p.locator(".rec-live")).to_contain_text("挪了没记")
    # 再改：一共少了 12.5
    p.get_by_label("微信实际余额").fill(str(round(wechat + 17.5, 2)))
    expect(p.locator(".rec-live")).to_contain_text("一共少了 ¥12.50")
    n = len(c.tx())
    p.get_by_role("button", name="对完了").click()
    c.wait_saved(n + 2)
    adj = c.tx()[-2:]
    assert all(t["type"] == "adjust" and t["rec"] == adj[0]["rec"] for t in adj), adj
    assert sum(t.get("spend", 0) for t in adj) == 12.5 and next(t for t in adj if t["account"] == "a-live").get("spend") == 12.5, adj
    assert c.data()["lastReconciled"] == TODAY
    p.wait_for_function("location.hash === '#/'")
    expect(p.locator(".health-tile", has_text="对账")).to_have_count(0)
    # 多了：不算收入，直接进存款
    c.go("#/reconcile")
    expect(p.locator(".reconcile-sub, .page-head")).to_contain_text("上次是")
    p.get_by_label("存钱卡实际余额").fill(str(round(c.balance("a-save") + 50, 2)))
    expect(p.locator(".rec-live")).to_contain_text("一共多了 ¥50")
    n = len(c.tx())
    p.get_by_role("button", name="对完了").click()
    c.wait_saved(n + 1)
    t = c.tx()[-1]
    assert t["amount"] == 50 and "spend" not in t, t


@step("总结：周（每天柱状图、环形图、和上周比、最大几笔）、月（曲线、每月存下、总资产、建议）")
def _(c):
    p = c.page
    c.go("#/summary")
    expect(p.get_by_role("img", name="每天的生活花销")).to_be_visible()
    expect(p.locator(".summary-head")).to_contain_text("这周")
    # 刊头：第几周宋体大字，右边日期范围和还剩几天；往前翻写「上周」
    expect(p.locator(".period-nav .mast-big")).to_have_text(f"第{date.today().isocalendar()[1]}周")
    expect(p.locator(".period-nav .mast-sub")).to_contain_text("这周")
    p.get_by_role("link", name="上一个").click()
    expect(p.locator(".period-nav .mast-sub")).to_have_text("上周")
    c.go("#/summary")
    # 这周没有「下一周」：占位的那个不显示（2026-10-25 样式被误删过，变成一个白圆圈）
    expect(p.locator(".period-nav .ghost")).to_be_hidden()
    expect(p.get_by_role("img", name="每天的生活花销")).to_contain_text("预算 ¥")
    expect(p.get_by_role("img", name="每天的生活花销")).not_to_contain_text("每天预算")
    p.screenshot(path=ART / "summary-week.png")
    # 最上面那张：大字花了多少、和计划比的标签和进度条、上周多少（先加两笔，看完还原）
    d0 = c.data()
    d = json.loads(json.dumps(d0))
    food = next(x["id"] for x in d["categories"] if x.get("group") == "food" and x["kind"] == "expense")
    acc = d["accounts"][0]["id"]
    monday = date.today() - timedelta(days=date.today().weekday())
    for i, (day, amt) in enumerate([(monday, 88), (monday - timedelta(days=3), 50)]):
        d["tx"].append({"id": f"wk{i}", "type": "expense", "date": day.isoformat(), "account": acc, "amount": amt, "category": food, "note": "", "createdAt": f"{day.isoformat()}T12:00:00Z"})
    try:
        c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
        c.go("#/summary")
        head = p.locator(".week-head")
        expect(head.locator(".big-num")).to_have_text(re.compile(r"^¥[\d,]+(\.\d\d)?$"))
        expect(head.locator(".wh-tag")).to_have_text(re.compile("节奏正常|花得有点快|下周收一收"))
        expect(head.locator(".time-legend")).to_contain_text("计划 ¥")
        expect(head.locator(".hero-foot")).to_contain_text("上周 ¥")
        head.screenshot(path=str(ART / "week-head.png"))
        # 月总结最上面也是一张卡：存下多少、收入 / 花了 / 其中生活、生活预算进度条、建议
        inc = next(x["id"] for x in d["categories"] if x["kind"] == "income")
        d["tx"].append({"id": "wk-in", "type": "income", "date": TODAY, "account": acc, "amount": 1000, "category": inc, "note": "", "createdAt": f"{TODAY}T12:00:00Z"})
        c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
        c.go("#/summary?mode=month")
        p.reload()  # 只换了 # 后面，要刷新才读到刚写的数据
        head = p.locator(".rep-head")
        expect(head.locator(".rep-title")).to_have_text(re.compile(r"^\d{4} 年 \d+ 月财务报表$"))
        expect(head.locator(".rep-title")).to_be_hidden()  # 网页上刊头写了，标题只印在 PDF 里
        expect(p.locator(".rep-kpi")).to_have_count(4)
        expect(head.locator(".advice")).to_be_hidden()  # 网页上只有四个数和图，小字只在 PDF 里
        expect(p.locator(".rep-kpi").first.locator("b")).to_have_text(re.compile(r"^¥[\d,]+(\.\d\d)?$"))
        # 储蓄率：大字 + 一句话 + 条（暑假那几个月没有收入时只写一句）
        meter = head.locator(".saving-meter")
        expect(meter).to_be_visible()
        if "没有收入" not in meter.inner_text():
            expect(meter.locator(".sm-num")).to_have_text(re.compile(r"^−?\d+%$"))
            expect(meter).to_contain_text(re.compile("每 100 块收入，存下 \\d+ 块|花的比收入多"))
            expect(meter.locator(".sm-mark.plan")).to_contain_text("按预算花")
        head.screenshot(path=str(ART / "month-head.png"))
    finally:
        c.data()  # 等网页把要上传的传完，再还原
        c.repo.external_write("finance.json", json.dumps(d0, ensure_ascii=False).encode())
        c.go("#/summary")
        p.reload()
    assert not any(t["id"].startswith("wk") for t in c.tx())
    if c.tx():
        expect(p.get_by_role("img", name="花在哪了")).to_be_visible()
    if p.locator(".legend-row").count():
        # 点一行：圆环里那一段突出，这一行亮一下，下面列出这一组的类别；再点一下收起
        row = p.locator(".legend-row").first
        total = p.locator(".donut-num").text_content()
        row.click()
        expect(row).to_have_attribute("aria-pressed", "true")
        expect(p.locator(".donut.has-sel .seg.on")).to_have_count(1)
        expect(p.locator(".donut-detail-box")).to_contain_text("里花在哪")
        expect(p.locator(".donut-sub")).to_contain_text("%")
        row.click()
        expect(p.locator(".donut-detail-box")).to_have_count(0)
        expect(p.locator(".donut-num")).to_have_text(total)
        # 点圆环的一段也行
        p.locator(".donut .seg").first.dispatch_event("click")
        expect(p.locator(".legend-row.on")).to_have_count(1)
    p.get_by_role("link", name="月", exact=True).click()
    # 月度报表：网页上只有四个关键数字和图；表格、小字都藏着，最下面是两个导出按钮
    expect(p.locator(".rep-sec h3").first).to_have_text("收支")
    for img in ("收支瀑布图", "月初和月底的钱按用途分", "每一块预算用了多少", "花钱曲线", "每天的生活花销日历", "收入、支出和储蓄率"):
        expect(p.get_by_role("img", name=img)).to_be_visible()
    expect(p.locator(".rep-table").first).to_be_attached()
    for el in p.locator(".rep-table").all():
        expect(el).to_be_hidden()
    assert p.evaluate("document.querySelector('.report').lastElementChild.classList.contains('rep-export')")
    expect(p.locator(".report")).not_to_contain_text("NaN")
    expect(p.locator(".report")).not_to_contain_text("undefined")
    # 对得上：按用途的合计 = 按账户的净资产（月初、现在都是）；心愿基金每一行进出加起来 = 月初到现在的变化
    got = p.evaluate("""() => Promise.all([import('./js/report.js'), import('./js/money.js')]).then(([r, m]) => {
      const d = JSON.parse(localStorage.getItem('ledger-cache') || 'null');
      const data = d && (d.data || d);
      return data && r.monthReport(data, m.ymd(new Date()), data.settings.usdRate || 7);
    })""")
    if got:
        for k in ("start", "end"):
            assert abs(sum(x[k] for x in got["purposes"]) - got[f"{k}NW"]["net"]) < 0.05, (k, got["purposes"], got[f"{k}NW"])
        assert abs(sum(x["v"] for x in got["spend"]) - got["key"]["total"]) < 0.01, got["spend"]
        w = got["notes"]["wish"]
        assert abs(w["start"] + sum(x["v"] for x in w["rows"]) - w["end"]) < 0.01, w
    expect(p.locator(".period-nav .mast-big")).to_contain_text("月")
    expect(p.locator(".period-nav .mast-sub")).to_contain_text("这个预算月")
    p.screenshot(path=ART / "report.png", full_page=True)
    # 存成 PDF：打开打印；打印时只留报表
    p.evaluate("window.print = () => { window.__printed = true; }")
    p.get_by_role("button", name="存成 PDF").click()
    assert p.evaluate("window.__printed") is True
    p.emulate_media(media="print")
    expect(p.locator("nav#nav")).to_be_hidden()
    expect(p.locator(".rep-export")).to_be_hidden()
    # PDF 里有完整的表格和明细
    for table in ("关键数字", "收支表", "按账户", "按用途", "预算执行", "花得最多的类别", "最近几个月"):
        expect(p.get_by_role("table", name=table)).to_be_visible()
    expect(p.get_by_role("table", name="按账户")).to_contain_text("净资产")
    expect(p.get_by_role("table", name="收支表")).to_contain_text("支出合计")
    expect(p.locator(".rep-head .advice")).to_contain_text("下个月可以试试")
    p.screenshot(path=ART / "report-print.png", full_page=True)
    p.emulate_media(media="screen")
    # 导出给 ChatGPT：下载成文件，八个部分都在
    p.get_by_role("button", name="导出给 ChatGPT").click()
    with p.expect_download() as dl:
        p.get_by_role("button", name="下载成文件").click()
    text = Path(dl.value.path()).read_text(encoding="utf-8")
    for want in ("财务报表", "## 一、关键数字", "## 二、收支表", "**收入合计**", "## 三、资产表", "**净资产**", "## 四、预算执行", "## 五、花销明细", "## 六、最近几个月", "## 七、附注", "不要推荐任何理财产品"):
        assert want in text, want
    p.locator(".sheet").get_by_role("button", name="取消").click()
    p.get_by_role("link", name="上一个").click()
    expect(p.locator(".rep-head")).to_be_visible()
    expect(p.locator(".period-nav")).to_be_visible()
    p.get_by_role("button", name="怎么用").click()
    expect(p.locator(".sheet")).to_contain_text("花钱曲线")
    expect(p.locator(".sheet")).to_contain_text("Save to Files")


def period_start(day, start=15):
    """预算月开始那天（和网页 periodOf 一样，startDay=15）"""
    if day.day >= start:
        return day.replace(day=start)
    first = day.replace(day=1) - timedelta(days=1)
    return first.replace(day=start)


@step("分类：老账本自动升级（三餐藏起来但老账照样显示）、日常先点小组再点类别、最近用过")
def _(c):
    p = c.page
    d = c.data()
    for x in d["categories"]:
        if x["id"] == "c-lunch":
            x.update(id="c-meal", name="三餐")  # 假装是老版本的账本
    d["categories"] = [x for x in d["categories"] if x["id"] not in ("c-breakfast", "c-dinner")]
    d.pop("categoryVersion", None)
    d["tx"].append({"id": "old1", "type": "expense", "date": TODAY, "account": "a-campus", "amount": 9, "category": "c-meal", "note": "老账", "createdAt": "2020-01-01T00:00:00Z"})
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    c.go("#/")
    p.reload()
    expect(c.recent("老账").first).to_contain_text("三餐")
    c.go("#/add")
    expect(p.get_by_role("button", name="早餐", exact=True)).to_be_visible()
    expect(p.get_by_role("button", name="三餐", exact=True)).to_have_count(0)
    expect(p.get_by_role("button", name="收纳整理")).to_have_count(0)
    p.get_by_role("group", name="日常小组").get_by_role("button", name="家居用品").click()
    p.get_by_role("button", name="收纳整理").click()
    p.get_by_label("金额", exact=True).fill("25")
    n = len(c.tx())
    p.get_by_role("button", name="记好了").click()
    c.wait_saved(n + 1)
    assert c.tx()[-1]["category"] == "c-storage"
    c.go("#/add")
    expect(p.get_by_role("group", name="最近用过的类别").get_by_role("button", name="收纳整理")).to_be_visible()
    # 「其他」：要写一下具体是什么，流水里显示写的名字
    p.get_by_role("group", name="日常小组").get_by_role("button", name="其他").click()
    p.get_by_role("group", name="其他类别").get_by_role("button", name="其他").click()
    expect(p.get_by_label("具体是什么")).to_be_focused()
    p.get_by_label("金额", exact=True).fill("680")
    n = len(c.tx())
    p.get_by_role("button", name="记好了").click()
    expect(p.locator(".toast.error")).to_contain_text("写一下具体是什么")
    p.get_by_label("具体是什么").fill("二手自行车")
    p.get_by_role("button", name="记好了").click()
    c.wait_saved(n + 1)
    t = c.tx()[-1]
    assert (t["category"], t["what"]) == ("c-other", "二手自行车"), t
    expect(c.recent("二手自行车").filter(has_text=re.compile(r"^二手自行车"))).to_contain_text("其他")


@step("心愿单：心愿基金（生活费结余）、大额要「开始攒」从 100 开始、排队、攒够才能买、不买了退回、冷静期、AI 建议、预下单")
def _(c):
    p = c.page
    c.go("#/more")
    expect(p.get_by_role("link", name="心愿单")).to_have_count(0)  # 心愿单藏在首页右上角
    c.go("#/")
    p.get_by_role("link", name="心愿单").click()
    expect(p.locator(".back-btn")).to_have_text("首页")
    for name, price, want in (("一本闲书", "80", "有了更好"), ("降噪耳机", "1200", "很想要"), ("机械键盘", "600", "有了更好")):
        p.get_by_role("button", name="加一个心愿").click()
        sheet = p.locator(".sheet")
        sheet.get_by_label("想要什么").fill(name)
        sheet.get_by_label("价格").fill(price)
        sheet.get_by_role("button", name=want, exact=True).click()
        if name == "降噪耳机":
            sheet.get_by_role("button", name="提升幸福感").click()
        sheet.get_by_role("button", name="加进心愿单").click()
        expect(p.locator(".wish", has_text=name)).to_be_visible()

    # 一张清单：一行一个，只写名字、价格和一句状态；点开是详情（弹出一页），按钮在详情里
    def open_wish(name):
        p.locator(".wish", has_text=name).click()
        return p.locator(".sheet")

    def close_sheet():
        p.locator(".sheet").get_by_role("button", name="关上").click()
        expect(p.locator(".sheet")).to_have_count(0)
    expect(p.locator(".wish", has_text="一本闲书").locator(".wish-status")).to_have_text("冷静中，还剩 3 天")
    expect(p.locator(".wish", has_text="降噪耳机")).not_to_contain_text("提升幸福感")
    expect(p.locator(".section-title", has_text="心愿")).to_have_count(0)  # 不再分小额、大额两段
    expect(p.locator(".wish-sort")).to_have_count(0)  # 没有排序按钮
    sheet = open_wish("降噪耳机")
    expect(sheet.locator(".wd-info")).to_contain_text("提升幸福感")
    expect(sheet.locator(".wd-info")).to_contain_text("很想要")
    close_sheet()
    assert next(w for w in c.data()["wishes"] if w["name"] == "降噪耳机")["kind"] == "joy"
    # 上个预算月：生活预算 3000，花了 2400 → 省下的进心愿基金；垫付那步报不回的 200 先从心愿基金扣（留 100）
    cur = period_start(date.today())
    prev = period_start(cur - timedelta(days=1))
    d = c.data()
    d["openingDate"] = prev.isoformat()
    d["tx"].append({"id": "lastpay", "type": "income", "date": prev.isoformat(), "account": "a-save", "amount": 8000, "category": "i-salary", "note": "上个月", "createdAt": "2020-01-01T00:00:00Z"})
    # 预算月头 7 天对的账，差额记在上个预算月最后一天（少了的算漏记的花销）；午饭少记这么多，上个月一共还是花 2400
    leak = sum(t.get("spend", 0) for t in d["tx"] if prev.isoformat() <= t["date"] < cur.isoformat() and t["type"] == "adjust")
    d["tx"].append({"id": "lastmonth", "type": "expense", "date": prev.isoformat(), "account": "a-live", "amount": 2400 - leak, "category": "c-lunch", "note": "上个月", "createdAt": "2020-01-01T00:00:00Z"})
    for w in d["wishes"]:
        if w["name"] == "降噪耳机":
            w["createdAt"] = prev.isoformat()
    living_before = sum(t["amount"] for t in d["tx"] if prev.isoformat() <= t["date"] < cur.isoformat() and t["type"] == "expense"
                        and t["category"] not in ("c-ai", "c-soft", "c-member", "c-fee", "c-trip", "c-wish"))
    living_before += leak
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    c.repo.external_write("config/ai.json", json.dumps({"deepseek": {"key": "sk-test", "model": "deepseek-flash"}}).encode())
    p.reload()
    base = round(3000 - living_before)
    writeoff = sum(t["amount"] for t in d["tx"] if t["type"] == "writeoff")
    fund = base - min(writeoff, max(0, base - 100))
    assert fund >= 260 and writeoff, (base, writeoff)
    expect(p.locator(".big-num")).to_have_text(f"¥{fund:,}")
    # 心愿基金只留一个数，来源点开才看
    expect(p.locator(".fund-rows")).to_have_count(0)
    p.locator(".wish-fund").get_by_role("button", name="来源 ›").click()
    expect(p.locator(".sheet .fund-rows")).to_contain_text(f"生活费省下+¥{base:,}")
    expect(p.locator(".sheet .fund-rows")).to_contain_text(f"帮专项和看病补−¥{base - fund:,}")
    close_sheet()
    # 大额心愿只在单子上，不攒钱，也不能买
    ear = p.locator(".wish", has_text="降噪耳机")
    expect(ear.locator(".wish-status")).to_have_text(re.compile("^还没开始攒"))
    sheet = open_wish("降噪耳机")
    expect(sheet.get_by_role("button", name="买了")).to_have_count(0)
    # 开始攒：从能随便用里拿 100
    sheet.get_by_role("button", name="开始攒").click()
    expect(ear.locator(".wish-status")).to_have_text(re.compile("^正在攒 · ¥100 / ¥1,200"))
    expect(ear.locator(".bar-track")).to_have_count(1)
    assert next(w for w in c.data()["wishes"] if w["name"] == "降噪耳机")["saveStart"] == TODAY
    expect(p.locator(".big-num")).to_have_text(f"¥{fund - 100:,}")  # 大字是现在能用的，攒给大额心愿的单独写
    p.locator(".wish-fund").get_by_role("button", name="来源 ›").click()
    expect(p.locator(".sheet .fund-split")).to_contain_text("另外 ¥100 攒给「降噪耳机」")
    expect(p.locator(".sheet .fund-rows")).to_contain_text("攒给大额心愿−¥100")
    close_sheet()
    # 一次只攒一个：键盘排队，可以不排
    kb = p.locator(".wish", has_text="机械键盘")
    open_wish("机械键盘").get_by_role("button", name="排队攒").click()
    expect(kb.locator(".wish-status")).to_have_text("排队第 1 个 · 冷静中，还剩 3 天")
    open_wish("机械键盘").get_by_role("button", name="不排了").click()
    expect(kb.locator(".wish-status")).to_have_text(re.compile("^还没开始攒"))
    assert "saveStart" not in next(w for w in c.data()["wishes"] if w["name"] == "机械键盘")
    # 排在哪按现在能做什么：正在攒的耳机在冷静中、排队的前面
    expect(p.locator(".wish .wish-name").first).to_have_text("降噪耳机")
    # 改成情怀、非常想要
    open_wish("机械键盘").get_by_role("button", name="改").click()
    p.locator(".sheet").get_by_role("button", name="非常想要").click()
    p.locator(".sheet").get_by_role("button", name="情怀").click()
    p.locator(".sheet").get_by_role("button", name="保存").click()
    expect(p.locator(".sheet")).to_have_count(0)
    expect(open_wish("机械键盘").locator(".wd-info")).to_contain_text("情怀")
    close_sheet()
    k = next(w for w in c.data()["wishes"] if w["name"] == "机械键盘")
    assert (k["want"], k["kind"]) == ("most", "feel"), k
    # AI：清单下面的小链接，弹出总建议和先后顺序；每个心愿的看法在它的详情里
    p.get_by_role("button", name="AI 看看 ›").click()
    p.locator(".sheet").get_by_role("button", name="问问 AI").click()
    expect(p.locator(".sheet .ai-summary")).to_contain_text("先买闲书")
    expect(p.locator(".sheet .wish-order li").first).to_have_text("一本闲书")
    close_sheet()
    expect(p.get_by_role("button", name=re.compile("^AI 建议 · "))).to_be_visible()
    sheet = open_wish("一本闲书")
    expect(sheet.locator(".wish-ai")).to_contain_text("想要")
    expect(sheet.locator(".wish-ai")).to_contain_text("AI 建议先后：第 1 个")
    close_sheet()
    assert "sk-test" not in json.dumps(c.data())  # 密钥不会写进账本
    sent = json.dumps(LAST_AI[0], ensure_ascii=False)
    assert "1/3" in sent and "开始攒" in sent, "大额心愿的规则要告诉 AI"
    # 单个「买了」的冷静期提醒
    open_wish("一本闲书").get_by_role("button", name="买了").click()
    expect(p.locator(".sheet")).to_contain_text("还在冷静期")
    p.locator(".sheet").get_by_role("button", name="取消").click()
    # 预下单：挑闲书和耳机，闲书改价 60
    small = fund - 100
    p.get_by_role("button", name="挑几个算一算 ›").click()
    p.get_by_role("checkbox", name="选上一本闲书").check()
    expect(p.locator(".cart-bar")).to_contain_text("选了 1 个 · 共 ¥80")
    p.get_by_role("checkbox", name="选上降噪耳机").check()
    expect(p.locator(".cart-bar")).to_contain_text("选了 2 个 · 共 ¥1,280")
    cart = lambda: [(x["id"], x.get("price")) for x in c.data()["wishCart"]]  # noqa: E731
    ids = {w["name"]: w["id"] for w in c.data()["wishes"]}
    assert cart() == [(ids["一本闲书"], None), (ids["降噪耳机"], None)], cart()
    price = p.get_by_label("一本闲书的价格")
    price.fill("60")
    price.press("Enter")
    price.blur()
    expect(p.locator(".cart")).to_contain_text("原价 ¥80")
    assert cart()[0] == (ids["一本闲书"], 60), cart()
    expect(p.locator(".cart")).to_contain_text(f"¥{small:,} → ¥{small - 60:,}")
    expect(p.locator(".cart")).to_contain_text("攒了 ¥100，还差 ¥1,100")
    expect(p.locator(".cart")).to_contain_text("没攒够的大额心愿先不能买")
    expect(p.locator(".cart")).to_contain_text("其中 1 个还在冷静期")
    expect(p.locator(".cart-bar")).to_contain_text(f"买完心愿基金剩 ¥{small - 60:,}")
    p.locator(".cart-bar").get_by_role("button", name="挑好了").click()
    expect(p.locator(".cart-bar")).to_have_count(0)
    expect(p.locator(".wish", has_text="一本闲书").locator(".wish-status")).to_contain_text("准备买")
    expect(p.get_by_role("button", name="预下单 · 2 个")).to_be_visible()
    # 都买了：耳机没攒够，不能勾
    n = len(c.tx())
    p.locator(".cart").get_by_role("button", name="都买了").click()
    sheet = p.locator(".sheet")
    expect(sheet.locator(".check-row", has_text="降噪耳机")).to_contain_text("还没攒够")
    expect(sheet.locator(".check-row", has_text="降噪耳机").get_by_role("checkbox")).to_be_disabled()
    expect(sheet).to_contain_text("共 ¥60")
    sheet.get_by_role("button", name="记好了").click()
    c.wait_saved(n + 1)
    t = c.tx()[-1]
    assert (t["category"], t["amount"], t["wishKind"], t["wish"]) == ("c-wish", 60, "small", ids["一本闲书"]), t
    book = next(w for w in c.data()["wishes"] if w["name"] == "一本闲书")
    assert (book["status"], book["boughtPrice"]) == ("bought", 60), book
    assert cart() == [(ids["降噪耳机"], None)], cart()
    expect(p.locator(".big-num")).to_have_text(f"¥{fund - 160:,}")
    p.locator(".cart").get_by_role("button", name="清空").click()
    expect(p.locator(".cart")).to_have_count(0)
    assert cart() == []
    # 再加一个 250 的小额心愿，能随便用的不够：差的算自由钱；攒着的不动
    n = len(c.tx())
    p.get_by_role("button", name="加一个心愿").click()
    p.locator(".sheet").get_by_label("想要什么").fill("台灯")
    p.locator(".sheet").get_by_label("价格").fill("250")
    p.locator(".sheet").get_by_role("button", name="加进心愿单").click()
    open_wish("台灯").get_by_role("button", name="买了").click()
    expect(p.locator(".sheet")).to_contain_text("差的")
    p.locator(".sheet").get_by_role("button", name="记好了").click()
    c.wait_saved(n + 2)
    split = sorted((x["category"], x["amount"]) for x in c.tx()[-2:])
    assert split == sorted([("c-wish", small - 60), ("c-like", 250 - (small - 60))]), split
    expect(p.locator(".big-num")).to_have_text("¥0")
    expect(p.locator(".wish", has_text="降噪耳机")).to_contain_text("¥100 / ¥1,200")
    # 不想要了：攒的退回能随便用
    open_wish("降噪耳机").get_by_role("button", name="不想要了").click()
    expect(p.locator(".toast")).to_contain_text("攒的 ¥100 退回心愿基金")
    expect(p.locator(".big-num")).to_have_text("¥100")
    open_wish("机械键盘").get_by_role("button", name="不想要了").click()
    # 实现了的、放下的合成最下面一行
    hist = p.locator(".wish-history")
    expect(hist.locator("summary")).to_have_text("实现了 2 个 · 放下 2 个")  # 放弃的不算省下多少钱
    # 放回来：大额心愿回到单子上，要重新点「开始攒」
    hist.locator("summary").click()
    expect(hist.locator(".wish-done")).to_contain_text("一本闲书")
    hist.locator(".wish-dropped .manage-row", has_text="降噪耳机").get_by_role("button", name="放回来").click()
    expect(p.locator(".wish", has_text="降噪耳机").locator(".wish-status")).to_have_text(re.compile("^还没开始攒"))
    expect(p.locator(".big-num")).to_have_text("¥100")
    assert "saveStart" not in next(w for w in c.data()["wishes"] if w["name"] == "降噪耳机")
    expect(hist.locator("summary")).to_have_text("实现了 2 个 · 放下 1 个")
    p.wait_for_timeout(600)
    p.screenshot(path=ART / "wishes.png", full_page=True)
    open_wish("降噪耳机")
    p.wait_for_timeout(400)
    p.screenshot(path=ART / "wish-sheet.png")
    close_sheet()
    # 「我们的花钱方式」里有心愿单和存款的规则
    c.go("#/rules")
    for t in ("心愿单：只花省下来的钱", "存款按用途分好"):
        p.locator(".rule-item summary", has_text=t).click()
    expect(p.locator(".rule-item", has_text="心愿单：只花省下来的钱")).to_contain_text("1/3")
    expect(p.locator(".rule-item", has_text="存款按用途分好")).to_contain_text("医疗备用金")

@step("买不买：先问问题再给建议、价格换算、硬规则（超过能自由用的钱一定不建议）、放进心愿单、决定买去记账、只发汇总、小课堂、回访")
def _(c):
    p = c.page
    c.go("#/")
    p.get_by_role("link", name="想买个东西？问问买不买……").click()
    p.wait_for_function("location.hash === '#/ask'")
    # 小课堂
    p.get_by_role("button", name=re.compile("^储蓄率")).click()
    expect(p.locator(".sheet")).to_contain_text("储蓄率 = 存下的钱 ÷ 收入")
    p.locator(".sheet").get_by_role("button", name="问问 AI").click()
    expect(p.get_by_label("想问什么")).to_have_value("关于「储蓄率」，我想问：")
    p.get_by_label("想问什么").fill("")
    p.get_by_label("想问什么").dispatch_event("input")
    # 先问问题
    p.get_by_role("button", name="想买个 1200 的机械键盘").click()
    expect(p.locator(".msg.ai").last).to_contain_text("现在用的键盘坏了吗")
    sent = json.dumps(LAST_AI[0], ensure_ascii=False)
    assert "食堂午饭" not in sent and "老账" not in sent, "流水明细不该发给 DeepSeek"
    assert "相当于" in sent and "天的饭钱" in sent and "冷静" in sent, "网站算好的换算和规则要发过去"
    assert "抽纸 | 清洁用品 | ×0" in sent and "有没有类似的" in sent, "物品档案里已有的东西要发过去（买前查重）"
    p.get_by_label("想问什么").fill("坏了，每天都用")
    p.get_by_role("button", name="发送").click()
    last = p.locator(".msg.ai").last
    expect(last.locator(".verdict")).to_have_text("等等再说")
    expect(last.locator(".facts")).to_contain_text("天的饭钱")
    expect(last.locator(".facts")).to_contain_text("先冷静")
    n = len(c.data().get("wishes", []))
    last.get_by_role("button", name="放进心愿单").click()
    expect(last).to_contain_text("你的决定：放进心愿单")
    d = c.data()
    assert len(d["wishes"]) == n + 1 and d["wishes"][-1]["name"] == "机械键盘" and d["decisions"][-1]["choice"] == "wish"
    # 硬规则：超过能自由用的钱，就算 AI 说可以也显示不建议
    p.get_by_label("想问什么").fill("想买个 99999 的电脑")
    p.get_by_role("button", name="发送").click()
    last = p.locator(".msg.ai").last
    expect(last.locator(".verdict")).to_have_text("不建议")
    expect(last.locator(".facts")).to_contain_text("能自由用的钱")
    last.get_by_role("button", name="决定买").click()
    p.wait_for_function("location.hash.startsWith('#/add')")
    expect(p.get_by_label("金额", exact=True)).to_have_value("99999")
    expect(p.get_by_label("备注")).to_have_value("顶配电脑")
    # 回访：一个多月前决定买的东西
    d = c.data()
    d["decisions"].append({"id": "old", "at": (date.today() - timedelta(days=40)).isoformat(), "item": "台灯", "price": 120, "verdict": "buy", "choice": "buy"})
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    c.go("#/ask")
    p.reload()
    expect(p.locator(".review")).to_contain_text("「台灯」买了一个多月了")
    p.locator(".review").get_by_role("button", name="值").click()
    expect(p.locator(".review")).to_have_count(0)
    assert next(x for x in c.data()["decisions"] if x["id"] == "old")["review"] == "worth"
    # 设置里有手机提醒
    c.go("#/settings")
    expect(p.locator(".card", has_text="手机提醒")).to_contain_text("每晚 9 点")

@step("个税退税：兼职收入填被预扣的个税；按年汇总；去年的点「办好了」记退税收入")
def _(c):
    p = c.page
    c.go("#/add")
    p.locator(".segmented").get_by_role("button", name="收入").click()
    p.get_by_label("金额", exact=True).fill("1600")
    p.get_by_role("button", name="兼职", exact=True).click()
    p.get_by_label("被预扣的个税").fill("400")
    p.get_by_role("group", name="账户").get_by_role("button", name="生活费卡").click()
    n = len(c.tx())
    p.get_by_role("button", name="记好了").click()
    c.wait_saved(n + 1)
    assert c.tx()[-1]["tax"] == 400
    # 去年的兼职（被扣过税）
    y = date.today().year - 1
    d = c.data()
    d["tx"].append({"id": "lastyearjob", "type": "income", "date": f"{y}-11-02", "account": "a-live", "amount": 800, "category": "i-job", "tax": 200, "note": "", "createdAt": "2020-01-01T00:00:00Z"})
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    c.go("#/more")
    p.reload()
    p.get_by_role("link", name=re.compile("^个税退税")).click()
    this = p.locator(".card").filter(has=p.get_by_role("heading", name=f"{y + 1} 年", exact=True))
    expect(this).to_contain_text("¥400")
    expect(this).to_contain_text("明年 3 月 1 日到 6 月 30 日办")
    last = p.locator(".card").filter(has=p.get_by_role("heading", name=f"{y} 年", exact=True))
    expect(last).to_contain_text("¥200")
    last.get_by_role("button", name="办好了").click()
    p.locator(".sheet").get_by_label("退了多少").fill("200")
    n = len(c.tx())
    p.locator(".sheet").get_by_role("button", name="办好了").click()
    c.wait_saved(n + 1)
    t = c.tx()[-1]
    assert (t["category"], t["amount"]) == ("i-tax", 200) and c.data()["taxYears"][str(y)]["refund"] == 200, t
    expect(last).to_contain_text("退了 ¥200")


@step("订阅：一年花多少、改金额、加年费、停掉、体检；满 3 个月首页提醒")
def _(c):
    p = c.page
    d = c.data()
    d["subReview"] = {"last": (date.today() - timedelta(days=100)).isoformat()}
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    c.go("#/")
    p.reload()
    expect(p.locator(".health-tile", has_text="订阅体检")).to_be_visible()
    c.go("#/subs")
    expect(p.locator(".banner")).to_contain_text("该体检了")
    card = p.locator(".card.sub", has_text="订阅乙")
    expect(card).to_contain_text("¥8,400/年")  # 100 美元 × 12 × 7
    card.get_by_role("button", name="考虑降档").click()
    expect(card.get_by_role("button", name="考虑降档")).to_have_attribute("aria-pressed", "true")
    card.get_by_role("button", name="改").click()
    p.locator(".sheet").get_by_label("每次扣多少").fill("20")
    p.locator(".sheet").get_by_role("button", name="保存").click()
    expect(card).to_contain_text("¥1,680/年")
    # 加一个年费
    p.get_by_role("button", name="加一个订阅").click()
    sheet = p.locator(".sheet")
    sheet.get_by_label("名称").fill("笔记软件")
    sheet.get_by_label("每次扣多少").fill("68")
    sheet.get_by_role("button", name="每年").click()
    sheet.get_by_label("每年哪天续费").fill(f"{date.today().year}-05-25")
    sheet.get_by_role("group", name="从哪个账户扣").get_by_role("button", name="生活费卡").click()
    sheet.get_by_role("button", name="保存").click()
    expect(p.locator(".card.sub", has_text="笔记软件")).to_contain_text("每年 5月25日")
    r = next(x for x in c.data()["recurring"] if x["name"] == "笔记软件")
    assert r["yearly"] == "05-25" and r["remindOnly"] and "day" not in r, r
    # 年费变成每年一次的专项（打开网站时生成），改金额跟着改，订阅页写一句不占订阅预算
    p.reload()
    p.wait_for_timeout(800)
    g = next(x for x in c.data()["goals"] if x.get("yearly") == r["id"])
    assert g["target"] == 68 and g["by"].endswith("-05-25") and "笔记软件" in g["name"], g
    expect(p.locator(".sub-month")).to_contain_text("在专项里提前留好")
    p.locator(".card.sub", has_text="笔记软件").get_by_role("button", name="改").click()
    p.locator(".sheet").get_by_label("每次扣多少").fill("90")
    p.locator(".sheet").get_by_role("button", name="保存").click()
    p.wait_for_timeout(500)
    assert next(x for x in c.data()["goals"] if x["id"] == g["id"])["target"] == 90
    p.locator(".card.sub", has_text="订阅甲").get_by_role("button", name="停掉", exact=True).click()
    expect(p.locator(".card.sub", has_text="订阅甲")).to_have_count(0)
    p.get_by_role("button", name="体检完了").click()
    expect(p.locator(".banner")).to_have_count(0)
    assert c.data()["subReview"]["last"] == TODAY
    c.go("#/")
    expect(p.locator(".health-tile", has_text="订阅体检")).to_have_count(0)


@step("专项存款：按日子排、每月要留多少、记账选专项不占预算、办完了结余回自由存款；钱都在哪")
def _(c):
    p = c.page
    c.go("#/goals")
    p.get_by_role("button", name="加一个专项").click()
    sheet = p.locator(".sheet")
    sheet.get_by_label("是什么事").fill("过渡金")
    sheet.get_by_label("一共要多少").fill("20000")
    sheet.get_by_label("哪天要用").fill("2030-12-31")
    sheet.get_by_role("button", name="保存").click()
    card = p.locator(".card.goal", has_text="过渡金")
    expect(card).to_contain_text("平均每月留")
    g = next(x for x in c.data()["goals"] if x["name"] == "过渡金")
    assert (g["target"], g["by"]) == (20000, "2030-12-31"), g
    # 再加一个近的：排到前面
    p.get_by_role("button", name="加一个专项").click()
    sheet.get_by_label("是什么事").fill("考研报名")
    sheet.get_by_label("一共要多少").fill("300")
    sheet.get_by_label("哪天要用").fill((date.today() + timedelta(days=40)).isoformat())
    sheet.get_by_role("button", name="保存").click()
    expect(p.locator(".card.goal").filter(has_text="考研报名")).to_contain_text("封存")
    goal_names = p.locator(".card.goal .wish-name").all_inner_texts()
    assert goal_names.index(next(x for x in goal_names if "考研报名" in x)) < goal_names.index(next(x for x in goal_names if "过渡金" in x)), goal_names
    # 记账时选「算在哪个专项里」：不占这个月预算
    n = len(c.tx())
    c.go("#/add")
    p.get_by_label("金额", exact=True).fill("320")
    p.get_by_role("button", name="耗材", exact=True).click()
    p.get_by_role("button", name="打印复印", exact=True).click()
    p.locator(".special-chips").get_by_role("button", name="考研报名").click()
    p.get_by_role("button", name="记好了").click()
    c.wait_saved(n + 1)
    t = c.tx()[-1]
    assert t["special"] == next(x for x in c.data()["goals"] if x["name"] == "考研报名")["id"], t
    c.go("#/goals")
    # 花够了（超了 ¥20）专项自己办完，不用再点；超的部分先从心愿基金扣
    expect(p.locator(".card.goal", has_text="考研报名")).to_have_count(0)
    p.locator(".done-wishes summary").click()
    expect(p.locator(".done-wishes")).to_contain_text("考研报名 ")
    expect(p.locator(".done-wishes")).to_contain_text("花了 ¥320 / ¥300")
    assert next(x for x in c.data()["goals"] if x["name"] == "考研报名")["status"] == "done"
    # 钱都在哪：圆圈和数额计算
    c.go("#/money")
    jar = lambda name: p.locator(".jar").filter(has=p.locator(".jar-name", has_text=re.compile(f"^{name}$")))  # noqa: E731
    group = lambda title: p.locator(".jar-group:not(.side), .jar-part").filter(has=p.locator(".jar-head", has_text=title))  # noqa: E731
    for name in ("家庭存款", "应急底线", "医疗备用金", "过渡金"):
        expect(group("存着不动").locator(".jar-name", has_text=name)).to_have_count(1)
    for name in ("自由存款", "心愿基金"):
        expect(group("能自由用").locator(".jar-name", has_text=name)).to_have_count(1)
    expect(group("这个月").locator(".jar-name")).to_have_text(["生活费还剩"])
    expect(p.locator(".jar-group.side .jar-head")).to_have_count(2)        # 能自由用和这个月并在一张卡里
    expect(p.locator(".card-check")).to_have_count(0)                    # 数额计算只在对账页
    # 点罐子：下面展开说明，再点收起
    jar("医疗备用金").click()
    expect(p.locator(".jar-detail")).to_contain_text("看病、买药从这里出")
    jar("医疗备用金").click()
    expect(p.locator(".jar-detail")).to_have_count(0)
    # 人情的专项只写人情和那个人
    for name in p.locator(".jar-name").all_inner_texts():
        assert not name.startswith("人情") or ("·" in name and "随礼" not in name), name
    c.go("#/reconcile")
    expect(p.locator(".card-check")).to_have_count(0)
    # 总结页可以整页翻到「钱都在哪」，再翻回来；离开总结页再回来默认是总结
    c.go("#/summary")
    p.get_by_role("button", name="翻到钱都在哪").click()
    expect(p.get_by_role("heading", name="钱都在哪")).to_be_visible()
    expect(jar("自由存款")).to_be_visible()
    p.get_by_role("button", name="翻回总结").click()
    expect(p.get_by_role("heading", name="总结")).to_be_visible()
    p.get_by_role("button", name="翻到钱都在哪").click()
    c.go("#/")
    c.go("#/summary")
    expect(p.get_by_role("heading", name="总结")).to_be_visible()

@step("没信号也能记：先存手机、立刻显示；离线刷新还在；有网后自动上传，和另一台设备的修改合并")
def _(c):
    p = c.page
    before = c.data()
    n = len(before["tx"])
    p.route(f"{API}/**", lambda r: r.abort())  # 断网
    c.go("#/add")
    p.get_by_label("金额", exact=True).fill("12")
    p.get_by_role("button", name="早餐", exact=True).first.click()
    p.get_by_label("备注").fill("离线记的")
    p.get_by_role("button", name="记好了").click()
    p.wait_for_function("location.hash === '#/'")
    expect(p.locator(".busy")).to_have_count(0)  # 不弹「正在保存」
    expect(p.locator(".tx", has_text="离线记的")).to_be_visible()
    expect(p.locator(".sync-pill")).to_contain_text("没网，1 项存在手机上")
    p.reload()
    expect(p.locator(".tx", has_text="离线记的")).to_be_visible()
    assert len(json.loads(c.repo.read("finance.json"))["tx"]) == n
    # 这期间另一台设备记了一笔
    d = json.loads(c.repo.read("finance.json"))
    d["tx"].append({"id": "otherdevice", "type": "expense", "date": TODAY, "account": "a-wechat", "amount": 7, "category": "c-drink", "note": "另一台设备", "createdAt": "2020-01-01T00:00:00Z"})
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    p.unroute(f"{API}/**")
    p.evaluate("window.dispatchEvent(new Event('online'))")
    d = c.data()
    notes = [t.get("note") for t in d["tx"]]
    assert "离线记的" in notes and "另一台设备" in notes and len(d["tx"]) == n + 2, notes[-3:]
    expect(p.locator(".sync-pill")).to_be_hidden()
    expect(c.recent("另一台设备")).to_have_count(1)
    expect(c.recent("离线记的")).to_have_count(1)

@step("暑假生活费自动目标、兼职三成进心愿基金、令牌到期提醒、预算调整建议（采用后改预算）、年度总结")
def _(c):
    p = c.page
    # 令牌还有 5 天过期
    exp = (date.today() + timedelta(days=5)).isoformat()
    p.evaluate(f"""() => {{ const s = JSON.parse(localStorage.getItem('inventory-settings')); s.tokenExpires = '{exp}'; localStorage.setItem('inventory-settings', JSON.stringify(s)); }}""")
    c.go("#/")
    p.reload()
    expect(p.locator(".token-banner")).to_contain_text("还有 5 天过期")
    # 暑假生活费
    c.go("#/goals")
    card = p.locator(".card.goal", has_text="暑假生活费")
    expect(card).to_contain_text("自动")
    expect(card.get_by_role("button", name="删掉")).to_have_count(0)
    # 兼职三成：前面记过 1600 + 去年 800 的兼职
    c.go("#/wishes")
    p.locator(".wish-fund").get_by_role("button", name="来源 ›").click()
    expect(p.locator(".sheet .fund-rows")).to_contain_text("兼职收入+¥720")
    p.locator(".sheet").get_by_role("button", name="关上").click()
    # 预算建议：挑最近 3 个完整、非暑假的预算月，日常每月只花 300
    d = c.data()
    cur = period_start(date.today())
    periods = []
    s0 = cur
    while len(periods) < 3:
        prev = period_start(s0 - timedelta(days=1))
        if prev.month not in (7, 8):
            periods.append((prev, s0))
        s0 = prev
    d["openingDate"] = period_start(s0 - timedelta(days=1)).isoformat()
    inside = lambda t: any(a.isoformat() <= t["date"] < b.isoformat() for a, b in periods)  # noqa: E731
    d["tx"] = [t for t in d["tx"] if not (inside(t) and t["type"] in ("expense", "writeoff"))]
    for a, _ in periods:
        d["tx"].append({"id": f"tissue{a}", "type": "expense", "date": a.isoformat(), "account": "a-live", "amount": 300, "category": "c-tissue", "note": "", "createdAt": "2020-01-01T00:00:00Z"})
        d["tx"].append({"id": f"lunch{a}", "type": "expense", "date": a.isoformat(), "account": "a-live", "amount": 1900, "category": "c-lunch", "note": "", "createdAt": "2020-01-01T00:00:00Z"})
        d["tx"].append({"id": f"fun{a}", "type": "expense", "date": a.isoformat(), "account": "a-live", "amount": 280, "category": "c-fun", "note": "", "createdAt": "2020-01-01T00:00:00Z"})
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    c.go("#/")
    p.reload()
    expect(p.locator(".cell", has_text="条调整建议")).to_be_visible()
    p.locator(".cell", has_text="条调整建议").click()
    expect(p.get_by_role("heading", name="花钱方式")).to_be_visible()
    adv = p.locator(".advice-card")
    daily = adv.locator(".advice-item", has_text="日常")
    expect(daily).to_contain_text("¥700 → ¥350")
    expect(daily).to_contain_text("纸巾清洁 月均 ¥300")
    expect(daily).to_contain_text("每月多存 ¥350")
    expect(adv.locator(".advice-item", has_text="订阅")).to_contain_text("¥900 → ¥150")
    expect(adv.locator(".advice-item", has_text="吃饭")).to_have_count(0)  # 吃饭花得和预算差不多，不用调
    adv.get_by_role("button", name=re.compile("^全部采用")).click()
    expect(adv).to_have_count(0)  # 都采用了：没有建议就不放这张卡
    expect(p.locator(".budget-plan")).to_contain_text("¥350")
    b = c.data()["budget"]
    assert (b["daily"], b["sub"], b["food"]) == (350, 150, 2000), b
    assert len(c.data()["budgetHistory"]) == 2
    # 年度报表：网页上四个数 + 图；表格只在 PDF 里；最下面三个导出
    year = date.today().year
    c.go("#/summary?mode=year")
    expect(p.locator(".period-nav .mast-big")).to_have_text(str(date.today().year))
    expect(p.locator(".period-nav .mast-sub")).to_contain_text("今年")
    p.locator(".period-nav").screenshot(path=ART / "mast-year.png")
    expect(p.locator(".rep-title")).to_have_text(f"{year} 年财务报表")
    expect(p.locator(".rep-title")).to_be_hidden()
    expect(p.locator(".rep-kpi")).to_have_count(4)
    for img in ("每月收入、支出和储蓄率", "存钱计划", "每一块每个月用了预算的多少", "钱花在哪了", "每个月花在哪", "这一年每天的生活花销"):
        expect(p.get_by_role("img", name=img)).to_be_visible()
    for el in p.locator(".rep-table").all():
        expect(el).to_be_hidden()
    days = (date(year, 12, 31) - date(year, 1, 1)).days + 1
    expect(p.locator(".year-cal rect")).to_have_count(days + 5 + 1)  # 每天一格 + 图例 5 格 + 今天描的一圈
    expect(p.locator(".report")).not_to_contain_text("NaN")
    meter = p.locator(".rep-head .saving-meter")
    expect(meter.locator(".sm-num")).to_have_text(re.compile(r"^−?\d+%$"))
    expect(meter.locator(".sm-mark.plan")).to_have_text(re.compile(r"^计划 \d+%$"))
    assert p.evaluate("document.querySelector('.report').lastElementChild.classList.contains('rep-export')")
    p.wait_for_timeout(600)  # 等淡入
    p.screenshot(path=ART / "year.png", full_page=True)
    # 一年的数对得上：每月加起来 = 全年
    y = p.evaluate("""() => Promise.all([import('./js/report.js'), import('./js/money.js')]).then(([r, m]) => {
      const d = JSON.parse(localStorage.getItem('ledger-cache')); const data = d.data || d;
      return r.yearReport(data, new Date().getFullYear(), data.settings.usdRate || 7);
    })""")
    act = [m for m in y["months"] if m["active"]]
    assert abs(sum(m["total"] for m in act) - y["key"]["total"]) < 0.05 and abs(sum(m["income"] for m in act) - y["key"]["income"]) < 0.05, act
    assert abs(sum(x["end"] for x in y["purposes"]) - y["endNW"]["net"]) < 0.05
    # PDF：表格都印出来
    p.emulate_media(media="print")
    for table in ("这一年", "每月收支", "按账户", "每月各块花了多少", "花得最多的类别", "附注"):
        expect(p.get_by_role("table", name=table)).to_be_visible()
    p.screenshot(path=ART / "year-print.png", full_page=True)
    p.emulate_media(media="screen")
    # Excel：总览、分表、每个月一张明细；数字是数字
    import openpyxl
    with p.expect_download() as dl:
        p.get_by_role("button", name="导出 Excel").click()
    assert dl.value.suggested_filename == f"财务报表-{year}.xlsx"
    dl.value.save_as(ART / "year.xlsx")
    wb = openpyxl.load_workbook(ART / "year.xlsx")
    assert wb.sheetnames[:7] == ["总览", "每月各块花销", "预算用了几成", "类别", "家底", "每天生活花销", "附注"], wb.sheetnames
    assert f"{date.today().month}月" in wb.sheetnames
    assert wb["总览"]["A2"].value == "收入" and isinstance(wb["总览"]["B2"].value, (int, float)), wb["总览"]["B2"].value
    month_sheet = wb[f"{date.today().month}月"]
    assert month_sheet["A1"].value == "日期" and month_sheet.max_row > 1
    # ChatGPT
    p.get_by_role("button", name="导出给 ChatGPT").click()
    with p.expect_download() as dl:
        p.get_by_role("button", name="下载成文件").click()
    text = Path(dl.value.path()).read_text(encoding="utf-8")
    for want in (f"# {year} 年财务报表", "## 每月收支", "## 每月用了预算的几成", "## 家底", "不要推荐任何理财产品"):
        assert want in text, want
    p.locator(".sheet").get_by_role("button", name="取消").click()
    # 月度报表也能导出 Excel
    c.go("#/summary?mode=month")
    with p.expect_download() as dl:
        p.get_by_role("button", name="导出 Excel").click()
    dl.value.save_as(ART / "month.xlsx")
    wb = openpyxl.load_workbook(ART / "month.xlsx")
    assert wb.sheetnames == ["总览", "家底", "预算", "类别", "每天生活花销", "明细"], wb.sheetnames
    # AI 年度点评：第二年 1 月打开去年的，自动写一次、存下来
    c.repo.external_write("config/ai.json", json.dumps({"deepseek": {"key": "sk-test", "model": "deepseek-flash"}}).encode())
    p.clock.set_fixed_time(f"{year + 1}-01-05T12:00:00")
    c.go(f"#/summary?mode=year&day={year}-12-31")
    p.reload()
    expect(p.locator(".rep-sec.letter .letter-text").first).to_be_visible(timeout=15000)
    assert str(year) in (c.data().get("letters") or {}), list((c.data().get("letters") or {}).keys())
    assert "一年结束了" in LAST_AI[0]["messages"][0]["content"]
    p.clock.set_fixed_time(date.today().isoformat() + "T12:00:00")
    c.go("#/")
    p.reload()

@step("最近记的搜索：全部时间里按名称、人名、金额、类别搜，显示合计")
def _(c):
    p = c.page
    expect(c.recent("二手")).to_have_count(1)
    expect(p.get_by_text("找到 1 笔 · 支出 ¥680")).to_be_visible()
    expect(c.recent("小乙").first).to_contain_text("小乙")
    expect(c.recent("没有这种东西")).to_have_count(0)
    expect(p.locator(".recent-list")).to_contain_text("没有找到")
    c.recent()


@step("发钱日一条龙：收入到了点一下记上，「还没」今天不再问，再点一下从存钱卡转生活费")
def _(c):
    p = c.page
    cur = period_start(date.today())
    if cur.month in (7, 8):
        return  # 暑假那两个月没有收入，这一步只在平时测
    d = c.data()
    d["openingDate"] = "2020-01-01"
    d["categories"] += [{"id": "i-a", "name": "收入甲", "kind": "income"}, {"id": "i-b", "name": "收入乙", "kind": "income"}]
    d["tx"] = [t for t in d["tx"] if not (t["date"] >= cur.isoformat() and t["type"] == "transfer" and t["account"] == "a-save")]
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    c.go("#/")
    p.reload()
    card = p.locator(".payday")
    expect(card).to_contain_text("收入甲的 ¥4,000 到了吗？")
    card.locator(".payday-row", has_text="收入乙").get_by_role("button", name="还没").click()
    expect(card.locator(".payday-row", has_text="收入乙")).to_have_count(0)
    card.locator(".payday-row", has_text="收入甲").get_by_role("button", name="到了").click()
    sheet = p.locator(".sheet")
    expect(sheet.get_by_label("金额")).to_have_value("4000")
    expect(sheet.get_by_role("group", name="到了哪个账户").get_by_role("button", name="存钱卡")).to_have_attribute("aria-pressed", "true")
    n = len(c.tx())
    sheet.get_by_role("button", name="记好了").click()
    c.wait_saved(n + 1)
    t = c.tx()[-1]
    assert (t["type"], t["category"], t["amount"], t["account"]) == ("income", "i-a", 4000, "a-save"), t
    expect(card).to_contain_text("挪到平时花钱的卡上")
    expect(card).to_contain_text("不影响算账")
    card.get_by_role("button", name="转到生活费卡").click()
    b = c.data()["budget"]
    living = b["food"] + b["daily"] + b["free"]
    expect(p.locator(".sheet").get_by_label("转多少")).to_have_value(str(living))
    p.locator(".sheet").get_by_role("button", name="转好了").click()
    c.wait_saved(n + 2)
    t = c.tx()[-1]
    assert (t["type"], t["account"], t["to"], t["amount"]) == ("transfer", "a-save", "a-live", living), t
    expect(p.locator(".payday")).to_have_count(0)

@step("钱到了没有：以前都是预算月头一天到的收入晚了 3 天还没记，琥珀色提一句")
def _(c):
    p = c.page
    cur = period_start(date.today())
    if cur.month in (7, 8):
        return
    keep = c.data()
    d = json.loads(json.dumps(keep))
    d["payday"] = {}
    olds, q = [], cur
    while len(olds) < 2:  # 前两个有收入的预算月（暑假不算）
        q = period_start(q - timedelta(days=1))
        if q.month not in (7, 8):
            olds.append(q)
    d["tx"] += [{"id": f"t-old-{i}", "type": "income", "date": day.isoformat(), "account": "a-save", "amount": 4000, "category": "i-b", "createdAt": "2020-01-01T00:00:00Z"}
                for i, day in enumerate(olds)]
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    c.go("#/")
    p.reload()
    row = p.locator(".payday-row", has_text="收入乙")
    late = (date.today() - cur).days
    if late >= 3:
        expect(row).to_contain_text(f"比平时晚了 {late} 天")
        expect(p.locator(".todo-item.warn", has=p.locator(".payday"))).to_have_count(1)
    else:
        expect(row).to_contain_text("到了吗？")
        expect(row).not_to_contain_text("晚了")
    c.repo.external_write("finance.json", json.dumps(keep, ensure_ascii=False).encode())
    p.reload()
    expect(p.locator(".spend-left")).to_be_visible()


@step("预测：首页月底大约花多少、超预算的可能；「往后」概率范围、底线、专项按时留够的可能、怎么算的；Chronos-2 结果混进来")
def _(c):
    p = c.page
    keep = c.data()
    d = json.loads(json.dumps(keep))
    y = date.today().year + 1
    d["yearly"] = [{"id": "fee", "name": "学费甲", "amount": 3000, "month": 3, "day": 1}]
    d["goals"] = [{"id": "g-far", "name": "过渡金甲", "target": 20000, "by": f"{y + 3}-05-31", "note": "", "createdAt": TODAY}]
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    c.go("#/")
    p.reload()
    line = p.locator(".spend-left .forecast-line")
    expect(line).to_contain_text("月底大约花")
    expect(line).to_contain_text("超预算的可能")
    line.click()
    p.wait_for_function("location.hash.includes('mode=ahead')")
    seg = p.locator(".segmented").first
    expect(seg.get_by_role("link", name="往后")).to_have_class(re.compile("on"))
    expect(p.locator(".mast-sub")).to_contain_text("次模拟 · 贝叶斯")
    expect(p.locator(".mast-sub")).not_to_contain_text("Chronos")
    expect(p.locator(".ahead-say")).to_contain_text("碰到底线的可能")
    expect(p.locator(".rep-kpi")).to_have_count(4)
    chart = p.locator(".rep-sec", has_text="每个月底的存款").locator("svg.chart")
    expect(chart).to_be_visible()
    expect(chart.locator("path[fill-opacity]")).to_have_count(2)  # 50%、80% 两层阴影
    big = p.locator(".rep-sec", has_text="要用大钱的时候")
    expect(big).to_contain_text("学费甲")
    expect(big).to_contain_text("暑假没有收入")
    expect(p.locator(".ahead-goal", has_text="过渡金甲")).to_contain_text("按时留够的可能")
    how = p.locator("details.how")
    how.locator("summary").click()
    for want in ("Dirichlet", "DP(α + n", "Beta(19.5, 0.5)", "Beta(4.5, 0.5)", "蒙特卡洛", "Chronos-2 还没有结果"):
        expect(how).to_contain_text(want)
    # 同一天、同一份账：再打开结果一样
    say = p.locator(".ahead-say").inner_text()
    p.reload()
    expect(p.locator(".ahead-say")).to_have_text(say)
    p.get_by_role("button", name="怎么用").click()
    expect(p.locator(".sheet")).to_contain_text("Chronos-2")
    p.keyboard.press("Escape")
    # Chronos-2 的结果（数据仓库 Actions 每天写）：混进来，「怎么算的」写最近一次
    levels = [0.01] + [round(0.05 * i, 2) for i in range(1, 20)] + [0.99]
    c.repo.external_write("config/forecast.json", json.dumps({"at": f"{TODAY}T01:00:00+00:00", "model": "amazon/chronos-2", "status": "ok", "from": TODAY,
        "days": 62, "context": 20, "levels": levels, "series": {"living": [[round(q * 150, 2) for q in levels] for _ in range(62)]}}).encode())
    c.repo.external_write("config/forecast-log.json", json.dumps({"levels": levels, "living": {TODAY: [round(q * 150, 2) for q in levels]}}).encode())
    p.reload()
    expect(p.locator(".mast-sub")).to_contain_text("贝叶斯 + Chronos-2")
    how.locator("summary").click()
    expect(how).to_contain_text("用了 20 天的记录")
    expect(how).to_contain_text("两个模型各占一半")
    p.wait_for_timeout(400)  # 等换页淡入完再截图
    p.screenshot(path=ART / "ahead.png", full_page=True)
    c.go("#/goals")
    expect(p.locator(".card.goal", has_text="过渡金甲")).to_contain_text("按时留够的可能")
    # 建议（第二批）：一个来不及的专项 → 「往后」有建议卡，首页「要做的事」提一句，点「做完了」这个预算月不再提
    soon = (date.today() + timedelta(days=150)).isoformat()
    d["goals"].append({"id": "g-hard", "name": "来不及乙", "target": 200000, "by": soon, "note": "", "createdAt": TODAY})
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    c.go("#/summary?mode=ahead")
    p.reload()
    advice = p.locator(".rep-sec", has_text="建议")
    expect(advice.locator(".advice-row.tone-warn").first).to_contain_text("可能", timeout=20000)  # 后台线程算，慢的机器要几秒
    expect(advice).to_contain_text("少花一半也到不了")
    p.wait_for_timeout(400)
    p.screenshot(path=ART / "ahead-advice.png", full_page=True)
    how.locator("summary").click()
    expect(how).to_contain_text("c* = min")
    c.go("#/")
    todo = p.locator(".todo-item", has_text="预测有建议")
    expect(todo).to_contain_text("还有", timeout=20000)  # 来不及乙、碰到底线……几条
    todo.get_by_role("button", name="做完了").click()
    expect(todo).to_have_count(0)
    assert "goal-g-hard" in c.data()["todoDone"]["forecast"], c.data()["todoDone"]
    c.repo.external_write("finance.json", json.dumps(keep, ensure_ascii=False).encode())
    c.go("#/")
    p.reload()
    expect(p.locator(".spend-left")).to_be_visible()


@step("记账时提示上次的价格；里程碑祝贺；月度小信（DeepSeek 写，存下来）")
def _(c):
    p = c.page
    # 上次的价格（「比平时贵」要有 3 次以上的记录）
    d = c.data()
    d["tx"].append({"id": "lunch3", "type": "expense", "date": max(date.today() - timedelta(days=1), period_start(date.today())).isoformat(), "account": "a-campus", "amount": 15, "category": "c-lunch", "note": "食堂午饭", "createdAt": "2020-01-01T00:00:00Z"})
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    c.go("#/")
    p.reload()
    expect(c.recent("食堂午饭").nth(2)).to_be_visible()  # 新数据读到了
    c.go("#/add")
    p.get_by_role("button", name="午餐", exact=True).first.click()
    expect(p.locator(".price-memo")).to_contain_text("上次")
    p.get_by_label("备注").fill("食堂午饭")
    expect(p.locator(".price-memo")).to_contain_text("最近 3 次平均 ¥15")
    p.get_by_label("金额", exact=True).fill("40")
    expect(p.locator(".price-memo")).to_contain_text("这次比平时贵一些")
    # 里程碑：收到一大笔钱，总资产跨过好几个整数
    d = c.data()
    d["tx"].append({"id": "bigincome", "type": "income", "date": TODAY, "account": "a-save", "amount": 100000, "category": "i-other", "note": "测试", "createdAt": "2020-01-01T00:00:00Z"})
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    c.go("#/")
    p.reload()
    card = p.locator(".milestone")
    expect(card).to_contain_text("总资产第一次超过 ¥100,000")
    reached = c.data()["milestones"]["reached"]
    assert "a100000" in reached and reached["a100000"]["at"] == TODAY, reached
    card.get_by_role("button", name="好").click()
    expect(p.locator(".milestone")).to_have_count(0)
    p.reload()
    expect(p.locator(".milestone")).to_have_count(0)
    # 月度小信：上个预算月的总结
    prev_end = period_start(date.today()) - timedelta(days=1)
    c.go(f"#/summary?mode=month&day={prev_end.isoformat()}")
    expect(p.locator(".letter-text")).to_contain_text("这个月你把吃饭控制得很好")
    letters = c.data()["letters"]
    assert any("吃饭控制得很好" in x["text"] for x in letters.values()), letters
    p.reload()
    expect(p.locator(".letter-text")).to_contain_text("这个月你把吃饭控制得很好")

@step("删一笔直接删、可以撤销")
def _(c):
    p = c.page
    c.add(18.5, cat="午餐", note="午饭")
    n = len(c.tx()) - 1
    t = c.tx()[-1]
    assert (t["amount"], t["category"], t["note"]) == (18.5, "c-lunch", "午饭"), t
    # 删掉刚才那笔：不弹确认，底部可以撤销
    c.go(f"#/add?edit={t['id']}")
    p.get_by_role("button", name="删除").click()
    toast = p.locator(".toast.undo")
    expect(toast).to_contain_text("删掉了：午餐")
    c.wait_saved(n)
    toast.get_by_role("button", name="撤销").click()
    expect(p.get_by_text("已撤销")).to_be_visible()
    c.wait_saved(n + 1)
    assert c.tx()[-1]["id"] == t["id"]


@step("账单查漏记：微信 Excel + 支付宝 GBK CSV；记过的跳过、退款提现不算、转账默认不记；补记带单号，记住付款方式；再导一次都对上")
def _(c):
    p = c.page
    c.add(15.03, cat="早餐", acc="微信", note="食堂")
    n = len(c.tx())
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in [["微信支付账单明细"], ["微信昵称：[测试]"], ["----------------------微信支付账单明细列表--------------------"],
                ["交易时间", "交易类型", "交易对方", "商品", "收/支", "金额(元)", "支付方式", "当前状态", "交易单号", "商户单号", "备注"],
                [f"{TODAY} 08:00:00", "商户消费", "食堂", "早餐", "支出", "¥15.03", "零钱", "支付成功", "wx1", "/", "/"],
                [f"{TODAY} 09:00:00", "商户消费", "瑞幸咖啡", "拿铁", "支出", "¥9.97", "零钱", "已全额退款", "wx2", "/", "/"],
                [f"{TODAY} 10:00:00", "零钱提现", "招商银行", "/", "/", "¥100.00", "零钱", "提现已到账", "wx3", "/", "/"],
                [f"{TODAY} 12:01:02", "商户消费", "美团外卖", "美团订单", "支出", "¥23.47", "零钱", "支付成功", "wx4", "/", "/"],
                [f"{TODAY} 18:00:00", "转账", "小王", "/", "支出", "¥50.01", "零钱", "对方已收钱", "wx5", "/", "/"]]:
        ws.append(row)
    wx = ART / "微信账单.xlsx"
    wb.save(wx)
    ali = ART / "支付宝账单.csv"
    ali.write_bytes(("支付宝交易明细\n交易时间,交易分类,交易对方,对方账号,商品说明,收/支,金额,收/付款方式,交易状态,交易订单号,商家订单号,备注,\n"
                     f"{TODAY} 20:00:00,交通出行,滴滴出行,x@x.com,打车,支出,18.61,中国银行储蓄卡(1234),交易成功,ali1\t,T1\t,,\n").encode("gbk"))
    c.go("#/more")
    p.get_by_role("link", name=re.compile("^账单查漏记")).click()
    p.get_by_label("选账单文件").set_input_files([str(wx), str(ali)])
    expect(p.get_by_text("没记的第 1 / 3 笔")).to_be_visible()
    expect(p.get_by_text("记过的 1 笔已经跳过")).to_be_visible()
    expect(p.get_by_text("美团外卖")).to_be_visible()
    expect(p.get_by_label("类别")).to_have_value("c-takeout")
    expect(p.get_by_role("group", name="账户").get_by_role("button", name="微信")).to_have_attribute("aria-pressed", "true")
    p.get_by_role("button", name="记上").click()
    expect(p.get_by_text("这是转给个人的钱")).to_be_visible()
    p.get_by_role("button", name="不记").click()
    expect(p.get_by_text("滴滴出行")).to_be_visible()
    expect(p.get_by_label("类别")).to_have_value("c-taxi")
    p.get_by_role("group", name="账户").get_by_role("button", name="生活费卡").click()
    p.get_by_role("button", name="记上").click()
    expect(p.get_by_text("补记 2 笔，共 ¥42.08")).to_be_visible()
    p.get_by_role("button", name="全部记上").click()
    expect(p.get_by_text("补记了 2 笔")).to_be_visible()
    d = c.data()
    new = d["tx"][n:]
    assert sorted((t["category"], t["amount"], t["account"], t["bill"]) for t in new) == [("c-takeout", 23.47, "a-wechat", "wx4"), ("c-taxi", 18.61, "a-live", "ali1")], new
    assert d["settings"]["payMethods"]["alipay:中国银行储蓄卡(1234)"] == "a-live"
    # 再导一次：都对上了
    p.get_by_role("button", name="再对一份").click()
    p.get_by_label("选账单文件").set_input_files([str(wx), str(ali)])
    expect(p.get_by_text("没有漏记的")).to_be_visible()


@step("导出全部账目 Excel")
def _(c):
    p = c.page
    c.go("#/more")
    with p.expect_download() as dl:
        p.get_by_role("button", name="导出全部账目（Excel）").click()
    path = ART / "ledger.xlsx"
    dl.value.save_as(path)
    import openpyxl
    rows = list(openpyxl.load_workbook(path).active.values)
    assert rows[0][0] == "日期" and len(rows) == len(c.tx()) + 1, rows[:2]


# 物品档案（买不买、心愿单建议会读）：抽纸用完了、购物清单上手动加了薯片和电池
def inventory_seed():
    I = lambda i, name, tag, **kw: {"id": i, "name": name, "assetId": None, "location": "L1", "tags": [tag], "quantity": 1, "fields": {},  # noqa: E731,E741
                                    "photos": [], "receipts": [], "notes": "", "archived": False, "consumable": False, **kw}
    data = {"version": 1, "tags": ["清洁用品", "电子产品"], "locations": [{"id": "L1", "name": "储物间", "parent": None, "assetId": None}],
            "items": [I("i1", "抽纸", "清洁用品", quantity=0, consumable=True, runningLow="2026-10-01"), I("i2", "台灯", "电子产品")],
            "shopping": {"extra": [{"id": "e1", "name": "薯片"}, {"id": "e2", "name": "电池"}], "skip": {}, "history": [], "toFile": []}}
    return {"inventory.json": json.dumps(data, ensure_ascii=False).encode()}


# 模拟交易的假行情：{ sym: { name, price, date, ask?, kind } }，测试里随时改
SIM = {
    "sh600001": {"name": "某某股份", "code": "600001", "price": 10.5, "date": "2026-10-26", "kind": "GP-A"},
    "sz159001": {"name": "某某纳指ETF", "code": "159001", "price": 2.0, "date": "2026-10-26", "kind": "QDII-ETF"},
    "usABC": {"name": "某美股", "code": "ABC.OQ", "price": 100.0, "date": "2026-10-26", "kind": "GP"},
    "sh510300": {"name": "沪深300ETF", "code": "510300", "price": 4.0, "date": "2026-10-26", "kind": "ETF"},
}


def sim_quote_line(sym):
    x = SIM[sym]
    f = ["0"] * 40
    f[1], f[2], f[3], f[4] = x["name"], x["code"], str(x["price"]), str(round(x["price"] / 1.01, 2))
    f[9], f[10] = str(round(x["price"] - 0.01, 2)), "100"
    f[19], f[20] = str(x.get("ask", x["price"])), "0" if x.get("ask") == 0 else "100"
    if x.get("ask") == 0:
        f[19] = "0"
    f[30] = x["date"] + " 16:00:01" if sym.startswith("us") else x["date"].replace("-", "") + "101500"
    f[31], f[32] = "0.10", "1.00"
    return f'v_{sym}="{"~".join(f)}";'


def fake_sim(page):
    def qt(route):
        syms = route.request.url.split("q=")[1].split(",")
        body = "\n".join(sim_quote_line(s) if s in SIM else 'v_pv_none_match="1";' for s in syms)
        route.fulfill(body=body.encode("gbk"), headers={"content-type": "text/html; charset=GBK", "access-control-allow-origin": "*"})

    def kline(route):
        k = route.request.url.split("param=")[1].split(",")[0]
        sym = k.split(".")[0]
        p = SIM.get(sym, {"price": 1})["price"]
        rows = [[f"2026-10-{d:02d}", str(p), str(p - 0.2 + d / 100), str(p), str(p), "1000"] for d in range(19, 27)]
        route.fulfill(json={"code": 0, "data": {k: {"qfqday": rows, "day": rows}}}, headers={"access-control-allow-origin": "*"})

    def smartbox(route):
        q = route.request.url.split("q=")[1]
        hits = [s for s, x in SIM.items() if quote(x["name"]) in q or x["code"].lower() in q.lower()]
        parts = [f'{s[:2]}~{SIM[s]["code"].lower() if s.startswith("us") else SIM[s]["code"]}~{SIM[s]["name"]}~py~{SIM[s]["kind"]}' for s in hits]
        route.fulfill(body=f'v_hint="{"^".join(parts)}"', headers={"content-type": "text/javascript; charset=utf-8"})

    page.route("https://qt.gtimg.cn/**", qt)
    page.route("https://web.ifzq.gtimg.cn/**", kline)
    page.route("https://smartbox.gtimg.cn/**", smartbox)


def fake_externals(page):
    page.route("https://api.frankfurter.dev/**", lambda r: r.fulfill(json={"base": "USD", "rates": {"CNY": 7.0}}))
    fake_sim(page)

    def deepseek(route):
        body = json.loads(route.request.post_data)
        user = body["messages"][-1]["content"]
        system = body["messages"][0]["content"]
        LAST_AI.clear()
        LAST_AI.append(body)
        if "模拟交易" in system:
            ans = {"questions": ["如果明天跌了 5%，你会怎么做？"]} if '"questions"' in system else {"text": "十月你买卖了一只股票，理由是「一直在涨」。\n想一想：涨了才买，是你一贯的想法吗？"}
            return route.fulfill(json={"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(ans, ensure_ascii=False)}}]})
        if "估一下" in system:
            ans = {"estimate": 150, "why": "编的估法"}
            return route.fulfill(json={"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(ans, ensure_ascii=False)}}]})
        if "理财小伙伴" in system:
            ans = {"letter": "这个月你把吃饭控制得很好，比预算少花了一些。下个月试试每周日看一眼周总结。"}
            return route.fulfill(json={"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(ans, ensure_ascii=False)}}]})
        if "理财小助手" in system:
            if "99999" in user:
                ans = {"answer": "这个太贵了，不过你想买的话也行。", "verdict": "buy", "item": {"name": "顶配电脑", "price": 99999}}
            elif "键盘" in user:
                ans = {"answer": "现在用的键盘坏了吗？多久用一次？", "verdict": None, "item": {"name": "机械键盘", "price": 1200}}
            elif "坏了" in user:
                ans = {"answer": "每天都用、旧的坏了，值得买；不过 1200 不少，先冷静几天、等双十一看看。", "verdict": "wait", "item": {"name": "机械键盘", "price": 1200}}
            else:
                ans = {"answer": "这个月吃饭花得最多，整体在预算内。", "verdict": None, "item": None}
            return route.fulfill(json={"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(ans, ensure_ascii=False)}}]})
        ids = {line.split(" | ")[1]: line.split(" | ")[0] for line in user.splitlines() if line.count(" | ") >= 5}
        ans = {"summary": "先买闲书，耳机等攒够再说。", "order": [ids.get("一本闲书"), ids.get("机械键盘"), ids.get("降噪耳机")],
               "items": [{"id": ids.get("一本闲书"), "when": "心愿基金够了，冷静期过了就买", "need": "想要", "comment": "想想会不会真的读完？"}]}
        route.fulfill(json={"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(ans, ensure_ascii=False)}}]})
    page.route("https://api.deepseek.com/**", deepseek)


@step("人情（不是钱的）：记一个、记礼物时对上就算还了、删了那笔又回到没还；节假日问这次还不还")
def _(c):
    p = c.page
    sheet = p.locator(".sheet")

    def favor(name, text, day):
        c.go("#/people")
        p.get_by_role("button", name="记一个人情").click()
        sheet.get_by_label("和谁", exact=True).fill(name[-1])
        sheet.get_by_role("option", name=name).click()
        expect(sheet.get_by_role("group", name="谁欠谁").get_by_role("button", name="我欠他")).to_have_attribute("aria-pressed", "true")
        sheet.get_by_label("什么事").fill(text)
        sheet.get_by_label("哪天").fill(day)
        n = len(c.data()["favors"])
        sheet.get_by_role("button", name="记好了").click()
        for _ in range(50):
            if len(c.data()["favors"]) == n + 1:
                break
            p.wait_for_timeout(200)
        return c.data()["favors"][-1]

    f1 = favor("小乙", "帮我搬书", "2026-09-01")
    assert f1["dir"] == "owe" and f1["status"] == "open" and f1["person"] == person(c, "小乙")["id"], f1
    expect(p.locator(".favor-list")).to_contain_text("小乙 · 帮我搬书")
    # 记一笔礼物，选上「还的是哪个人情」
    c.go("#/add")
    p.get_by_label("金额", exact=True).fill("88")
    p.get_by_role("button", name="人情社交", exact=True).click()
    p.get_by_role("button", name="礼物", exact=True).first.click()
    p.get_by_role("group", name="还的人情").get_by_role("button", name="小乙 · 帮我搬书").click()
    n = len(c.tx())
    p.get_by_role("button", name="记好了").click()
    c.wait_saved(n + 1)
    t = c.tx()[-1]
    f = next(x for x in c.data()["favors"] if x["id"] == f1["id"])
    assert t["favor"] == f1["id"] and f["status"] == "done" and f["doneTx"] == t["id"], (t, f)
    c.go(f"#/person/{f1['person']}")
    expect(p.locator(".favor-row.done")).to_contain_text("礼物 · 给小乙 ¥88")
    # 删掉这笔：人情回到没还
    c.recent("礼物").first.click()
    p.get_by_role("button", name="删除").click()
    c.wait_saved(n)
    f = next(x for x in c.data()["favors"] if x["id"] == f1["id"])
    assert f["status"] == "open" and "doneTx" not in f, f
    f2 = favor("小甲", "请我喝咖啡", "2026-09-02")
    # 国庆前一天：问这次还不还
    p.clock.set_fixed_time("2026-09-30T12:00:00")
    c.go("#/")
    card = p.locator(".favor-holiday")
    expect(card).to_contain_text("明天开始放国庆了")
    card.locator(".favor-ask", has_text="帮我搬书").get_by_role("button", name="这次还").click()
    card.locator(".favor-ask", has_text="请我喝咖啡").get_by_role("button", name="这次不还").click()
    expect(card).not_to_contain_text("请我喝咖啡")
    card.screenshot(path=ART / "favor-holiday.png")
    fs = {x["id"]: x for x in c.data()["favors"]}
    assert fs[f1["id"]]["plan"] == "2026-国庆" and fs[f2["id"]]["skip"] == "2026-国庆", fs
    # 这次要还的：点「记一笔」，聚餐请客和人情都先选好了
    card.get_by_role("link", name="记一笔").click()
    expect(p.get_by_role("group", name="还的人情").get_by_role("button", name="小乙 · 帮我搬书")).to_have_attribute("aria-pressed", "true")
    expect(p.locator(".picker-token")).to_have_text("小乙×")
    p.get_by_label("金额", exact=True).fill("120")
    p.get_by_role("group", name="账户").get_by_role("button", name="微信", exact=True).click()
    p.get_by_role("button", name="记好了").click()
    c.wait_saved(n + 1)
    assert c.tx()[-1]["category"] == "c-social" and c.tx()[-1]["who"] == f1["person"]
    assert next(x for x in c.data()["favors"] if x["id"] == f1["id"])["status"] == "done"
    # 平常的周末不问；下个假期（元旦前一天）再问「这次不还」的那个
    p.clock.set_fixed_time("2026-11-14T12:00:00")
    c.go("#/")
    expect(p.locator(".favor-holiday")).to_have_count(0)
    p.clock.set_fixed_time("2026-12-31T12:00:00")
    c.go("#/")
    expect(p.locator(".favor-holiday")).to_contain_text("请我喝咖啡")
    # 没花钱的还法
    c.go("#/people")
    p.locator(".favor-row", has_text="请我喝咖啡").get_by_role("button", name="还了").click()
    sheet.get_by_label("怎么还的").fill("帮他占了座")
    sheet.get_by_role("button", name="还了").click()
    expect(p.locator(".favor-row", has_text="请我喝咖啡")).to_have_count(0)
    f = next(x for x in c.data()["favors"] if x["id"] == f2["id"])
    assert f["status"] == "done" and f["doneNote"] == "帮他占了座", f
    p.clock.set_fixed_time(date.today().isoformat() + "T12:00:00")
    # 要准备多少钱：DeepSeek 估一个；300 以上自动变成存款目标；买不买也算进去
    c.repo.external_write("config/ai.json", json.dumps({"deepseek": {"key": "sk-test", "model": "deepseek-flash"}}).encode())
    c.go("#/people")
    p.reload()
    p.get_by_role("button", name="记一个人情").click()
    sheet.get_by_label("和谁", exact=True).fill("乙")
    sheet.get_by_role("option", name="小乙").click()
    sheet.get_by_label("什么事").fill("请我吃火锅")
    sheet.get_by_label("他为我花了多少").fill("120")
    sheet.get_by_role("button", name="估一个").click()
    expect(sheet.get_by_label("预计要准备多少")).to_have_value("150")
    expect(sheet.get_by_text("AI：编的估法")).to_be_visible()
    assert "他为我花了大约 ¥120" in json.dumps(LAST_AI[-1], ensure_ascii=False)
    sheet.get_by_role("button", name="记好了").click()
    for _ in range(50):
        if any(x["text"] == "请我吃火锅" for x in c.data()["favors"]):
            break
        p.wait_for_timeout(200)
    f3 = next(x for x in c.data()["favors"] if x["text"] == "请我吃火锅")
    assert f3["cost"] == 120 and f3["estimate"] == 150, f3
    p.get_by_role("button", name="记一个人情").click()
    sheet.get_by_label("和谁", exact=True).fill("甲")
    sheet.get_by_role("option", name="小甲").click()
    sheet.get_by_label("什么事").fill("婚礼随礼")
    sheet.get_by_label("预计要准备多少").fill("500")
    sheet.get_by_role("button", name="记好了").click()
    plan = p.locator(".social-plan")
    expect(plan).to_contain_text("日常里出")
    expect(plan).to_contain_text("小乙 · 请我吃火锅")
    expect(plan).to_contain_text("大额，专项存款")
    c.go("#/goals")
    expect(p.locator(".card.goal", has_text="人情：小甲 · 婚礼随礼")).to_contain_text("¥500")
    c.go("#/")
    # 首页只列 45 天内要准备的：没写日子的人情按下一个假期算，离得远就不在首页
    soon = p.evaluate("""async () => {
      const { nextHoliday } = await import('./js/cal.js');
      const { today, } = await import('./js/util.js');
      const t = today(); const h = nextHoliday(t); const due = !h || h.start < t ? t : h.start;
      const lim = new Date(); lim.setDate(lim.getDate() + 45);
      return due <= lim.toISOString().slice(0, 10);
    }""")
    if soon:
        expect(p.locator(".social-plan")).to_contain_text("¥150")
    else:
        expect(p.locator(".social-plan", has_text="请我吃火锅")).to_have_count(0)
    c.go("#/ask")
    p.get_by_label("想问什么").fill("想买个 100 的台灯")
    p.get_by_role("button", name="发送").click()
    for _ in range(50):
        if LAST_AI and "理财小助手" in LAST_AI[-1]["messages"][0]["content"]:
            break
        p.wait_for_timeout(200)
    sent = json.dumps(LAST_AI[-1], ensure_ascii=False)
    assert "要准备的人情钱" in sent and "婚礼随礼" in sent and "人情：小甲 · 婚礼随礼" in sent, sent[:3000]
    # 礼尚往来：收到的红包礼金，选谁给的
    c.go("#/add")
    p.locator(".segmented").get_by_role("button", name="收入", exact=True).click()
    p.get_by_label("金额", exact=True).fill("66")
    p.get_by_role("group", name="收入来源").get_by_role("button", name="收到的红包礼金").click()
    p.get_by_label("谁给的").click()
    expect(p.locator(".picker-cap")).to_have_text("最近")
    p.get_by_label("谁给的").fill("甲")
    p.get_by_role("option", name="小甲").click()
    n = len(c.tx())
    p.get_by_role("button", name="记好了").click()
    c.wait_saved(n + 1)
    assert c.tx()[-1]["who"] == f2["person"] and c.tx()[-1]["category"] == "i-gift"
    c.go(f"#/person/{f2['person']}")
    expect(p.locator(".tx-list").first).to_contain_text("收到的红包礼金 · 小甲给的")


@step("首页：还能花在最上面，要做的事合成一张卡，单独超支的组不说「一切正常」，健康指标不留半行")
def _(c):
    p = c.page
    c.go("#/")
    expect(p.locator(".hero")).to_be_visible()
    # 「还能花」在「要做的事」前面
    assert p.evaluate("() => { const t = document.querySelector('.todo'); return !t || !!(document.querySelector('.hero').compareDocumentPosition(t) & Node.DOCUMENT_POSITION_FOLLOWING); }")
    if p.locator(".todo").count():
        expect(p.locator(".todo .todo-head h3")).to_have_text("要做的事")
        expect(p.locator(".todo-count")).to_contain_text("件")
        # 每个要留意 / 要处理的指标单独一行（以前只露第一个）
        if p.locator(".todo-more").count():
            p.locator(".todo-more").click()
        for name in p.locator(".health-tile.warn .tile-name, .health-tile.bad .tile-name").all_inner_texts():
            expect(p.locator(".todo-item", has_text=name)).to_have_count(1)
    # 「做完了」：对账点了记成今天对过，不再提醒；撤销回来
    d = c.data()
    d["lastReconciled"] = (date.today() - timedelta(days=30)).isoformat()
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    p.reload()
    p.locator(".todo-more").click() if p.locator(".todo-more").count() else None
    p.get_by_role("button", name="对账做完了").click()
    expect(p.locator(".todo-item", has_text="对账")).to_have_count(0)
    assert c.data()["lastReconciled"] == date.today().isoformat()
    p.get_by_role("button", name="撤销").click()
    expect(p.locator(".todo-item", has_text="对账")).to_have_count(1)
    # 别的指标点了只是这个预算月不提（记号 = 预算月开始日 | 轻重）
    tiles = p.locator(".todo-item.warn .todo-done, .todo-item.bad .todo-done")
    if tiles.count() > 1:
        name = tiles.nth(0).get_attribute("aria-label")
        tiles.nth(0).click()
        expect(p.get_by_role("button", name=name)).to_have_count(0)
        p.get_by_role("button", name="撤销").click()
        expect(p.get_by_role("button", name=name)).to_have_count(1)
    # 首页不再放「钱都在哪」的入口（总结页翻过来就是它）
    expect(p.locator("#view a[href='#/money']")).to_have_count(0)
    # 即将扣款单独一张卡：一笔一行，美元账户最下面写够不够；健康指标里没有它了
    expect(p.locator(".health-tile", has_text="即将扣款")).to_have_count(0)
    expect(p.locator(".charges .charge-row").first).to_be_visible()
    expect(p.locator(".charges .charge-foot")).to_contain_text("Apple ID余额")
    # 年费快到了：只在即将扣款卡里，不再占一个健康指标格子、也不进要做的事
    d = c.data()
    d["recurring"].append({"id": "r-yearly-test", "name": "编的年费", "amount": 50, "account": "a-live", "yearly": (date.today() + timedelta(days=10)).strftime("%m-%d"), "remindOnly": True, "note": "到期前想想还续不续"})
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    p.reload()
    expect(p.locator(".charges .charge-row", has_text="编的年费")).to_contain_text("到期前想想还续不续")
    expect(p.locator(".health-tile", has_text="编的年费")).to_have_count(0)
    expect(p.locator(".todo-item", has_text="编的年费")).to_have_count(0)
    # 只看 14 天内：20 天后的不出现
    d["recurring"][-1]["yearly"] = (date.today() + timedelta(days=20)).strftime("%m-%d")
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    p.reload()
    expect(p.locator(".charges .charge-row", has_text="编的年费")).to_have_count(0)
    d["recurring"] = [r for r in d["recurring"] if r["id"] != "r-yearly-test"]
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    p.reload()
    # 一切正常那句只在没有要留意的指标时出现（前面几步花了多少、今天几号会让花钱节奏有时偏快）
    warns = p.locator(".health-tile.warn, .health-tile.bad").count()
    expect(p.locator(".summary-title")).to_have_count(0 if warns else 1)
    # 健康指标两列：窄的格子是双数
    assert p.locator(".health-tile:not(.wide)").count() % 2 == 0
    hd = p.evaluate("""() => import('./js/money.js').then((m) => m.headline({ items: [{ key: 'pace', level: 'good', ratio: 1 }], overGroups: [{ name: '日常', over: 540 }] }))""")
    assert hd["title"] == "整体还在计划里" and "日常超了 ¥540" in hd["text"], hd
    hd = p.evaluate("""() => import('./js/money.js').then((m) => m.headline({ items: [{ key: 'pace', level: 'good', ratio: 1 }], overGroups: [] }))""")
    assert hd["title"].startswith("一切正常"), hd


@step("记账时用心愿基金的钱（不占预算、从能随便用的扣，不够不让记）；订阅页写清这个月的订阅钱")
def _(c):
    p = c.page
    c.add(100, kind="收入", cat="兼职", acc="存钱卡")   # 兼职三成进心愿基金，保证有钱可用
    c.go("#/wishes")
    have = float(p.locator(".big-num").inner_text().replace("¥", "").replace(",", ""))
    assert have >= 30, have
    c.go("#/add")
    p.get_by_label("金额", exact=True).fill(str(int(have) + 50))
    p.get_by_role("button", name="饮料奶茶", exact=True).first.click()
    p.get_by_role("group", name="账户").get_by_role("button", name="生活费卡", exact=True).click()
    sw = p.get_by_role("checkbox", name="用心愿基金的钱")
    expect(p.locator(".fund-hint")).to_contain_text(f"心愿基金能用 ¥{int(have):,}")
    sw.check()
    expect(p.locator(".fund-hint")).to_contain_text("不够这一笔")
    n = len(c.tx())
    p.get_by_role("button", name="记好了").click()
    expect(p.locator(".toast")).to_contain_text("不够这一笔")
    assert len(c.tx()) == n
    p.get_by_label("金额", exact=True).fill("20")
    expect(p.locator(".fund-hint")).to_contain_text(f"记好后心愿基金还能用 ¥{int(have) - 20:,}")
    p.get_by_role("button", name="记好了").click()
    c.wait_saved(n + 1)
    t = c.tx()[-1]
    assert t["wishFund"] is True and t["category"] == "c-drink" and t["amount"] == 20, t
    # 首页和周总结写清是心愿基金的钱（不说成手续费、出差自付），首页进度条和周总结一样：紫色是花了多少
    c.go("#/")
    expect(p.locator(".legend")).to_contain_text("心愿基金")
    expect(p.get_by_text(re.compile(r"另有 ¥[\d,.]+ (是心愿基金的钱，不占预算|不占预算：.*心愿基金)"))).to_be_visible()
    expect(p.locator(".spend-left .time-fill.wh-fill")).to_be_visible()
    c.go("#/summary?mode=week")
    expect(p.locator(".week-head .from-saved")).to_contain_text("心愿基金")
    c.go("#/wishes")
    expect(p.locator(".big-num")).to_have_text(f"¥{int(have) - 20:,}")
    c.recent(t["date"])
    expect(p.locator(f'.recent-list .tx[href="#/add?edit={t['id']}"]')).to_contain_text("心愿基金")  # 按 id 找：网页时钟被前面固定过，同一天的先后不可靠
    # 改这一笔、去掉勾：回到普通花销
    c.go(f"#/add?edit={t['id']}")
    expect(p.get_by_role("checkbox", name="用心愿基金的钱")).to_be_checked()
    p.get_by_role("checkbox", name="用心愿基金的钱").uncheck()
    p.get_by_role("button", name="保存").click()
    p.wait_for_timeout(500)
    assert "wishFund" not in next(x for x in c.tx() if x["id"] == t["id"])
    # 订阅页：这个预算月的订阅（已扣、还要扣、够不够预算），美元订阅的钱怎么走
    c.go("#/subs")
    card = p.locator(".sub-month")
    expect(card).to_contain_text("这个预算月的订阅")
    expect(card).to_contain_text("预算 ¥")
    expect(card.locator(".sub-sum")).to_contain_text("都扣完大约")
    expect(p.locator(".sub-flow")).to_contain_text("记一笔「转账」")
    expect(p.locator(".sub-flow .charge-foot")).to_contain_text("Apple ID余额")


@step("专项「办完了」：跳到记账，专项、金额、账户、类别先填好，点记好了才扣；花够了自己办完")
def _(c):
    p = c.page
    d = c.data()
    d["goals"].append({"id": "g-fee", "name": "2099 年编的费用", "target": 800, "by": "2099-10-20", "category": "c-special", "createdAt": TODAY})
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    p.reload()
    c.go("#/goals")
    p.locator(".card.goal", has_text="编的费用").get_by_role("button", name="办完了").click()
    expect(p.get_by_label("金额", exact=True)).to_have_value("800")
    expect(p.locator(".special-chips .chip.on")).to_contain_text("编的费用")
    assert next(g for g in c.data()["goals"] if g["id"] == "g-fee").get("status") is None  # 还没记，不算办完
    n = len(c.tx())
    expect(p.get_by_role("button", name="专项", exact=True)).to_have_class(re.compile("on"))  # 类别里有「专项」，办完了带过来
    p.get_by_role("button", name="记好了").click()
    c.wait_saved(n + 1)
    assert next(g for g in c.data()["goals"] if g["id"] == "g-fee").get("status") == "done"  # 花够了自己办完
    t = c.tx()[-1]
    assert t["special"] == "g-fee" and t["amount"] == 800 and t["category"] == "c-special", t


@step("外观和「生活」一致：问候、预算月和发工资小标签、还能花的圆环、角落一句；记账页不放")
def _(c):
    p = c.page
    c.go("#/")
    expect(p.locator(".today-head .greet")).to_have_text(re.compile("好|夜深"))
    expect(p.locator(".head-tags .tag").first).to_have_text(re.compile("月预算 · 第 \\d+ 天"))
    expect(p.locator(".spend-left .ring")).to_be_visible()
    expect(p.locator(".hero .big-num")).to_have_text(re.compile(r"^¥[\d,]+\.\d\d$"))  # 精确到分
    expect(p.locator(".hero .time-legend")).to_contain_text("天")
    expect(p.locator(".health-tile").first.locator(".ring.mini")).to_be_visible()
    if p.locator(".health-tile", has_text="即将扣款").count():
        expect(p.locator(".health-tile", has_text="即将扣款")).to_have_class(re.compile("wide"))
    expect(p.locator(".whisper")).to_have_count(1)
    p.wait_for_timeout(600)
    p.screenshot(path=ART / "look-home.png", full_page=True)
    c.go("#/add")
    expect(p.locator(".whisper")).to_have_count(0)
    c.go("#/summary")   # 角落那句话只在首页
    expect(p.locator(".whisper")).to_have_count(0)


@step("模拟交易：开始、搜一只、按真的规则买（一手、写理由、AI 问一句）、当天不能卖、看着的、第二天卖、美股、回头看、月度复盘、导出、重开一局")
def _(c):
    p = c.page
    c.repo.external_write("config/ai.json", json.dumps({"deepseek": {"key": "sk-test", "model": "deepseek-flash"}}).encode())
    p.clock.set_fixed_time("2026-10-26T10:00:00+08:00")  # 周一，A 股开着
    c.go("#/sim")
    p.reload()
    expect(p.locator("nav#nav")).to_contain_text("模拟")
    p.get_by_role("button", name="开始").click()
    expect(p.locator(".sim-worth .big-num")).to_have_text("¥10,000")
    expect(p.locator(".sim-markets")).to_contain_text("A 股开盘中")
    # 对照：10/25 收盘 4.05 全买沪深 300，现在 4.00
    expect(p.locator(".sim-bench")).to_contain_text("一直拿沪深 300：¥9,876.54 · 你比它多 ¥123.46")
    # 搜一只、打开
    p.get_by_role("button", name="搜一只").click()
    p.get_by_label("搜股票").fill("600001")
    p.locator(".sim-results .cell", has_text="某某股份").click()
    expect(p.locator(".sim-price")).to_have_text("¥10.50")
    expect(p.locator(".sim-quote")).to_contain_text("今天买的明天才能卖")
    # 买：不到一手不行，要写为什么、选拿多久，AI 问一句
    p.get_by_role("button", name="买", exact=True).click()
    sheet = p.locator(".sheet")
    sheet.get_by_label("买多少股").fill("50")
    sheet.get_by_label("为什么买").fill("一直在涨")
    sheet.get_by_role("button", name="几周").click()
    sheet.get_by_role("button", name="下一步").click()
    expect(sheet).to_contain_text("最少买 100 股")
    sheet.get_by_label("买多少股").fill("200")
    sheet.get_by_role("button", name="下一步").click()
    expect(sheet).to_contain_text("点一下有几成把握")
    sheet.get_by_role("button", name="把握 3").click()
    sheet.get_by_label("买多少股").fill("200")
    expect(sheet).to_contain_text("手续费 ¥5.02")
    sheet.get_by_role("button", name="下一步").click()
    sheet.get_by_label("如果明天跌了 5%，你会怎么做？").fill("再等等")
    sheet.get_by_role("button", name="确认买").click()
    expect(p.locator(".sheet")).to_have_count(0)
    t = c.data()["simTrades"][0]
    assert (t["side"], t["qty"], t["price"], t["fee"], t["cny"], t["hold"], t["sure"]) == ("buy", 200, 10.5, 5.02, 2105.02, "weeks", 3), t
    assert "把握：3/5" in LAST_AI[0]["messages"][-1]["content"]
    expect(p.locator(".sim-quote .chart-mark")).to_have_count(1)
    expect(p.locator(".sim-quote")).to_contain_text("我买的那天")
    assert t["ask"] == [{"q": "如果明天跌了 5%，你会怎么做？", "a": "再等等"}], t["ask"]
    assert "模拟交易" not in json.dumps(LAST_AI[0]["messages"][-1]) or "一直在涨" in LAST_AI[0]["messages"][-1]["content"]
    # 今天买的今天不能卖
    p.get_by_role("button", name="卖", exact=True).click()
    sheet.get_by_label("为什么卖").fill("试试")
    sheet.get_by_role("button", name="下一步").click()
    expect(sheet).to_contain_text("今天买的明天才能卖")
    sheet.get_by_role("button", name="取消").click()
    # 看着的
    c.go("#/sim/s/sz159001")
    expect(p.locator(".sim-quote")).to_contain_text("当天买当天能卖")
    p.get_by_role("button", name="加到看着的").click()
    p.locator(".sheet").get_by_label("为什么关注").fill("想看看")
    p.get_by_role("button", name="加上").click()
    c.go("#/sim")
    expect(p.locator(".group", has_text="某某纳指ETF")).to_contain_text("想看看")
    # 返回：左上角「‹ 模拟」回到模拟页，还在原来滚到的地方
    p.evaluate("window.scrollTo(0, 400)")
    y = p.evaluate("window.scrollY")
    assert y > 0, y
    p.locator(".sim-row", has_text="某某纳指ETF").click()
    expect(p.locator(".sim-price")).to_have_text("¥2.000")
    back = p.locator(".back-btn")
    expect(back).to_have_text("模拟")
    back.click()
    expect(p).to_have_url(URL + "#/sim")
    expect(p.locator(".sim-worth")).to_be_visible()
    p.wait_for_function(f"window.scrollY === {y}")
    # 直接打开股票页（没有上一页）也能回到模拟
    p.evaluate("sessionStorage.clear()")
    p.goto(URL + "#/sim/s/sh600001")
    p.reload()
    p.evaluate("history.replaceState(null, '')")
    p.reload()
    p.locator(".back-btn").click()
    expect(p).to_have_url(URL + "#/sim")
    expect(p.locator(".sim-worth")).to_contain_text("现金 ¥7,894.98")
    # 第二天卖出
    p.clock.set_fixed_time("2026-10-27T10:00:00+08:00")
    SIM["sh600001"].update(price=11.0, date="2026-10-27")
    c.go("#/sim/s/sh600001")
    p.reload()
    p.get_by_role("button", name="卖", exact=True).click()
    expect(sheet.get_by_label("卖多少股")).to_have_value("200")
    sheet.get_by_label("为什么卖").fill("涨够了")
    sheet.get_by_role("button", name="把握 4").click()
    sheet.get_by_role("button", name="下一步").click()
    sheet.get_by_role("button", name="确认卖").click()
    expect(p.locator(".sheet")).to_have_count(0)
    sell = c.data()["simTrades"][1]
    assert (sell["side"], sell["price"], sell["fee"]) == ("sell", 10.99, 6.12), sell  # 买一价 10.99；佣金 5 + 过户费 0.02 + 印花税 1.10
    expect(p.locator(".sim-quote .chart-mark")).to_have_count(2)
    # 美股：晚上开盘，按汇率换
    p.clock.set_fixed_time("2026-10-27T22:00:00+08:00")
    SIM["usABC"]["date"] = "2026-10-27"
    c.go("#/sim/s/usABC")
    p.reload()
    p.get_by_role("button", name="买", exact=True).click()
    sheet.get_by_label("买多少股").fill("1")
    sheet.get_by_label("为什么买").fill("想试试美股")
    sheet.get_by_role("button", name="还不知道").click()
    sheet.get_by_role("button", name="把握 1").click()
    sheet.get_by_role("button", name="下一步").click()
    sheet.get_by_role("button", name="确认买").click()
    expect(p.locator(".sheet")).to_have_count(0)
    p.screenshot(path=ART / "sim-stock.png", full_page=True)
    us = c.data()["simTrades"][2]
    assert (us["rate"], us["cny"]) == (7, round((100 + 1.99) * 7, 2)), us
    # 一周后：回头看那笔卖出；上个月有买卖 → 自动写十月的复盘
    p.clock.set_fixed_time("2026-11-03T10:00:00+08:00")
    c.go("#/sim")
    p.reload()
    look = p.locator(".sim-look")
    expect(look).to_contain_text("涨够了")
    look.get_by_role("button", name="想对了").click()
    look.get_by_label("回头看想说的").fill("卖得刚好")
    look.get_by_role("button", name="记下").click()
    expect(p.locator(".sim-review")).to_contain_text("2026 年 10 月的复盘")
    d = c.data()
    assert d["simTrades"][1]["look"]["verdict"] == "right" and d["simTrades"][1]["look"]["text"] == "卖得刚好"
    assert [r["month"] for r in d["simReviews"]] == ["2026-10"]
    p.screenshot(path=ART / "sim.png", full_page=True)
    # 导出：下载成文件，里面有理由、AI 问答、回头看、复盘
    p.get_by_role("button", name="导出给 ChatGPT").click()
    with p.expect_download() as dl:
        p.get_by_role("button", name="下载成文件").click()
    text = Path(dl.value.path()).read_text(encoding="utf-8")
    for want in ("不要直接推荐", "理由：一直在涨；打算拿 几周；把握 3/5", "一直拿沪深 300", "对照：如果开局时", "AI 问：如果明天跌了 5%", "回头看", "想对了，卖得刚好", "在观察的", "2026 年 10 月的复盘"):
        assert want in text, want
    p.locator(".sheet").get_by_role("button", name="复制，去 ChatGPT").evaluate("b => b.closest('.sheet-overlay').remove()")
    # 模拟的钱不进账本
    assert "某某股份" not in json.dumps(c.tx(), ensure_ascii=False)
    # 重开一局
    c.go("#/sim")
    p.get_by_role("button", name="重开一局").click()
    expect(p.locator(".sim-worth .big-num")).to_have_text("¥10,000")
    expect(p.locator(".group", has_text="笔买卖")).to_contain_text("3 笔买卖")
    games = c.data()["simGames"]
    assert len(games) == 2 and games[0]["end"] == "2026-11-03" and "end" not in games[1], games



@step("返回：子页面左上角回到上一页、写着上一页的名字；刷新、直接打开也行；底栏的几页没有返回")
def _(c):
    p = c.page
    c.go("#/more")
    expect(p.locator(".back-btn")).to_have_count(0)
    p.locator(".cell", has_text="账户").first.click()
    back = p.locator(".back-btn")
    expect(back).to_have_text("更多")
    p.locator(".cell", has_text="存钱卡").first.click()
    expect(back).to_have_text("账户")
    p.screenshot(path=ART / "back.png")
    back.click()
    expect(p).to_have_url(URL + "#/accounts")
    expect(back).to_have_text("更多")
    back.click()
    expect(p).to_have_url(URL + "#/more")
    expect(p.locator(".back-btn")).to_have_count(0)
    # 刷新以后返回还认得上一页
    p.locator(".cell", has_text="设置").click()
    p.reload()
    expect(p.locator(".back-btn")).to_have_text("更多")
    # 直接打开子页面：回到它上面那一页
    p.evaluate("sessionStorage.clear(); history.replaceState(null, '', '#/budget')")
    p.reload()
    expect(p.locator(".back-btn")).to_have_text("更多")
    p.locator(".back-btn").click()
    expect(p).to_have_url(URL + "#/more")
    # 改一笔：从首页点进去，返回回首页
    d = c.data()
    d["tx"].append({"id": "t-back", "type": "expense", "date": TODAY, "account": "a-live", "amount": 12, "category": d["categories"][0]["id"], "note": "返回测试", "createdAt": TODAY + "T12:00:00Z"})
    c.repo.external_write("finance.json", json.dumps(d, ensure_ascii=False).encode())
    c.go("#/")
    p.reload()
    c.recent("返回测试").first.click()  # 同一天记的多，最近 8 笔里不一定有它，用搜索找
    expect(p.locator(".page-head h1")).to_have_text("改一笔")
    expect(p.locator(".back-btn")).to_have_text("首页")
    p.screenshot(path=ART / "back-edit.png")
    p.locator(".back-btn").click()
    expect(p).to_have_url(URL + "#/")
    for tab in ("#/", "#/summary", "#/add"):
        c.go(tab)
        expect(p.locator(".page-head").first).to_be_visible()
        expect(p.locator(".back-btn")).to_have_count(0)


def main():
    only = sys.argv[1:]
    ART.mkdir(exist_ok=True)
    repo = FakeRepo({"README.md": b"# finance-data\n"})
    inventory = FakeRepo(inventory_seed())
    serve({REPO: repo, "x/inventory-data": inventory}, API_PORT)
    handler = partial(SimpleHTTPRequestHandler, directory=str(ROOT))
    handler.log_message = lambda *a: None
    # 浏览器一次要并发拉二十多个脚本文件；默认只排 5 个等待连接，偶尔会丢掉请求、整页加载不出来
    ThreadingHTTPServer.request_queue_size = 128
    app = ThreadingHTTPServer(("127.0.0.1", APP_PORT), handler)
    threading.Thread(target=app.serve_forever, daemon=True).start()

    failed, errors, ran = [], [], 0
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        ctx = browser.new_context(viewport={"width": 390, "height": 844}, accept_downloads=True, locale="zh-CN")
        ctx.set_default_timeout(10000)
        page = ctx.new_page()
        page.on("pageerror", lambda e: errors.append(str(e)))
        page.on("dialog", lambda d: d.accept(c.prompt) if d.type == "prompt" else d.accept())
        c = Ctx(page, repo, inventory)
        fake_externals(page)
        for i, (name, fn) in enumerate(STEPS):
            if i and only and not any(k in name for k in only):
                continue
            ran += 1
            try:
                fn(c)
                print(f"  ✓ {name}")
            except Exception:
                failed.append(name)
                page.screenshot(path=ART / f"fail-{i:02d}.png", full_page=True)
                print(f"  ✗ {name}\n{traceback.format_exc()}")
                if i == 0:
                    break
        browser.close()
    if errors:
        print("页面报错：", *errors, sep="\n  ")
    print(f"\n{ran - len(failed)}/{ran} 通过" + ("" if ran == len(STEPS) else f"（共 {len(STEPS)} 项，没跑完）"))
    sys.exit(1 if failed or errors else 0)


if __name__ == "__main__":
    main()
